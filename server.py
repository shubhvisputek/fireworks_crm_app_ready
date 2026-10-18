#!/usr/bin/env python3
import os, json, csv, io, re, sqlite3, hashlib, hmac, base64, secrets, zipfile, datetime as dt, urllib.parse, traceback, math
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from email.parser import BytesParser
from email.policy import default
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import load_workbook, Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.getenv('FIREWORKS_DB', os.path.join(BASE_DIR, 'data', 'fireworks.db'))
STATIC_DIR = os.path.join(BASE_DIR, 'static')
EXPORT_DIR = os.path.join(BASE_DIR, 'exports')
BILLS_DIR = os.path.join(BASE_DIR, 'bills')
SECRET_PATH = os.getenv('FIREWORKS_SECRET', os.path.join(BASE_DIR, 'data', '.secret'))
PORT = int(os.getenv('PORT', '8000'))
HOST = os.getenv('HOST', '127.0.0.1')
COOKIE = 'fw_session'
MAX_BODY_BYTES = int(os.getenv('FIREWORKS_MAX_BODY_BYTES', str(10 * 1024 * 1024)))
LOGIN_MAX_FAILURES = 10
LOGIN_WINDOW_SECONDS = 15 * 60
SECURITY_HEADERS = {
    'X-Content-Type-Options': 'nosniff',
    'X-Frame-Options': 'DENY',
    'Referrer-Policy': 'same-origin',
    # Inline handlers are used throughout the UI, so scripts still allow 'unsafe-inline'.
    'Content-Security-Policy': "default-src 'self'; script-src 'self' 'unsafe-inline'; style-src 'self' 'unsafe-inline'; img-src 'self' data: blob:; object-src 'none'; base-uri 'none'; form-action 'self'; frame-ancestors 'none'",
}

os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
os.makedirs(EXPORT_DIR, exist_ok=True)
os.makedirs(BILLS_DIR, exist_ok=True)


def now_iso():
    return dt.datetime.now().astimezone().isoformat(timespec='seconds')


def today():
    return dt.date.today().isoformat()


def money(x):
    return float(Decimal(str(x or 0)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))


def parse_bool(value):
    if isinstance(value, bool):
        return value
    return str(value or '').strip().lower() in ('1','true','yes','on')


def allocate_discount(discount, line_totals):
    total_cents = int(Decimal(str(money(discount))) * 100)
    weights = [int(Decimal(str(money(value))) * 100) for value in line_totals]
    weight_total = sum(weights)
    if not total_cents or not weight_total:
        return [0.0 for _ in weights]
    shares = [(total_cents * weight) // weight_total for weight in weights]
    remainder = total_cents - sum(shares)
    order = sorted(range(len(weights)), key=lambda i: (total_cents * weights[i]) % weight_total, reverse=True)
    for index in order[:remainder]:
        shares[index] += 1
    return [share / 100 for share in shares]


def is_loopback_host(host):
    return host in ('127.0.0.1', 'localhost', '::1')


def clean_date(value, field, required=False):
    value=str(value or '').strip()
    if not value:
        if required: raise ValueError(f'{field} is required.')
        return None
    try:
        return dt.date.fromisoformat(value).isoformat()
    except ValueError:
        raise ValueError(f'{field} must be a valid date (YYYY-MM-DD).')


def excel_safe(value):
    # Stop text such as "=HYPERLINK(...)" being written as a live spreadsheet formula.
    if isinstance(value, str) and value[:1] in ('=', '+', '-', '@', '\t', '\r'):
        return "'" + value
    return value


_login_failures={}

def login_blocked(username):
    now=dt.datetime.now().timestamp()
    attempts=[t for t in _login_failures.get(username,[]) if now-t < LOGIN_WINDOW_SECONDS]
    _login_failures[username]=attempts
    return len(attempts) >= LOGIN_MAX_FAILURES


def record_login_failure(username):
    if len(_login_failures)>10000: _login_failures.clear()
    _login_failures.setdefault(username,[]).append(dt.datetime.now().timestamp())


def json_response(obj, status=200):
    return status, json.dumps(obj, default=str).encode('utf-8')


def sign_token(payload):
    raw = base64.urlsafe_b64encode(json.dumps(payload, separators=(',', ':')).encode()).decode().rstrip('=')
    sig = hmac.new(get_secret(), raw.encode(), hashlib.sha256).hexdigest()
    return raw + '.' + sig


def verify_token(token):
    try:
        raw, sig = token.split('.', 1)
        good = hmac.new(get_secret(), raw.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(sig, good): return None
        payload = json.loads(base64.urlsafe_b64decode(raw + '=' * (-len(raw) % 4)))
        if payload.get('exp', 0) < dt.datetime.now().timestamp(): return None
        return payload
    except Exception:
        return None


def get_secret():
    if not os.path.exists(SECRET_PATH):
        with open(SECRET_PATH, 'wb') as f:
            f.write(secrets.token_bytes(32))
    return open(SECRET_PATH, 'rb').read()


def hash_password(password, salt=None):
    salt = salt or secrets.token_bytes(16)
    rounds = 120000
    dk = hashlib.pbkdf2_hmac('sha256', password.encode(), salt, rounds)
    return f'{rounds}${base64.b64encode(salt).decode()}${base64.b64encode(dk).decode()}'


def safe_user(u):
    if not u: return None
    d=dict(u)
    d.pop('password_hash',None); d.pop('session_version',None)
    return d


def is_admin_user(user):
    return bool(user and user.get('role') == 'OWNER')


def staff_api_allowed(method, path):
    if method == 'GET':
        return bool(
            path in {
                '/api/dashboard', '/api/analytics', '/api/expenses',
                '/api/products', '/api/held-bills', '/api/backup',
                '/api/export/sales.xlsx', '/api/export/sales-detailed.xlsx',
                '/api/export/sales.pdf'
            }
            or re.fullmatch(r'/api/held-bills/resume/\d+', path)
            or re.fullmatch(r'/api/invoice-pdf/\d+', path)
        )
    if method == 'POST':
        return bool(
            path in {
                '/api/sales', '/api/expenses', '/api/recurring-expenses',
                '/api/held-bills', '/api/password'
            }
            or re.fullmatch(r'/api/held-bills/resume/\d+', path)
        )
    return False


def check_password(password, stored):
    try:
        rounds_s, salt_b64, hash_b64 = stored.split('$')
        dk = hashlib.pbkdf2_hmac('sha256', password.encode(), base64.b64decode(salt_b64), int(rounds_s))
        return hmac.compare_digest(base64.b64encode(dk).decode(), hash_b64)
    except Exception: return False


def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA foreign_keys = ON')
    conn.execute('PRAGMA journal_mode = WAL')
    conn.execute('PRAGMA busy_timeout = 5000')
    return conn


def exec_script(conn, script):
    conn.executescript(script)
    conn.commit()

SCHEMA = '''
CREATE TABLE IF NOT EXISTS users (
  id INTEGER PRIMARY KEY, username TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL,
  role TEXT NOT NULL DEFAULT 'OWNER', active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS business_settings (
  id INTEGER PRIMARY KEY CHECK(id=1), data_json TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS products (
  id INTEGER PRIMARY KEY, sku TEXT UNIQUE NOT NULL, name TEXT NOT NULL, category TEXT, subcategory TEXT,
  unit TEXT NOT NULL DEFAULT 'piece', units_per_box REAL NOT NULL DEFAULT 1, boxes_per_carton REAL NOT NULL DEFAULT 1,
  units_per_carton REAL NOT NULL DEFAULT 1, brand TEXT, hsn TEXT, gst_rate REAL NOT NULL DEFAULT 0,
  cess_rate REAL NOT NULL DEFAULT 0, reorder_level REAL NOT NULL DEFAULT 0, status TEXT NOT NULL DEFAULT 'ACTIVE',
    cost_price REAL NOT NULL DEFAULT 0, mrp REAL NOT NULL DEFAULT 0,
    retail_selling_price REAL NOT NULL DEFAULT 0, wholesale_selling_price REAL NOT NULL DEFAULT 0,
    default_unit_price REAL NOT NULL DEFAULT 0,
  barcode TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS price_versions (
  id INTEGER PRIMARY KEY, product_id INTEGER NOT NULL REFERENCES products(id), effective_date TEXT NOT NULL,
  mrp REAL NOT NULL DEFAULT 0, retail_price REAL NOT NULL DEFAULT 0, wholesale_price REAL NOT NULL DEFAULT 0,
  cost_price REAL NOT NULL DEFAULT 0, created_by INTEGER REFERENCES users(id), created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_price_product_date ON price_versions(product_id, effective_date, id);
CREATE TABLE IF NOT EXISTS stock_batches (
  id INTEGER PRIMARY KEY, product_id INTEGER NOT NULL REFERENCES products(id), batch_number TEXT, supplier TEXT,
  manufacturer TEXT, manufacture_date TEXT, purchase_date TEXT, qty_base REAL NOT NULL DEFAULT 0,
  cost REAL NOT NULL DEFAULT 0, location TEXT, notes TEXT, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS inventory_movements (
  id INTEGER PRIMARY KEY, product_id INTEGER NOT NULL REFERENCES products(id), movement_type TEXT NOT NULL,
  qty_base REAL NOT NULL, unit_cost REAL, ref_type TEXT, ref_id INTEGER, reason TEXT, user_id INTEGER,
  batch_id INTEGER REFERENCES stock_batches(id), created_at TEXT NOT NULL, notes TEXT
);
CREATE INDEX IF NOT EXISTS idx_inv_product_date ON inventory_movements(product_id, created_at);
CREATE TABLE IF NOT EXISTS suppliers (
  id INTEGER PRIMARY KEY, name TEXT NOT NULL, phone TEXT, gstin TEXT, address TEXT, notes TEXT, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS purchases (
  id INTEGER PRIMARY KEY, purchase_number TEXT UNIQUE NOT NULL, purchase_date TEXT NOT NULL, supplier_id INTEGER REFERENCES suppliers(id),
  supplier_name TEXT, invoice_number TEXT, subtotal REAL NOT NULL DEFAULT 0, gst REAL NOT NULL DEFAULT 0,
  freight REAL NOT NULL DEFAULT 0, total REAL NOT NULL DEFAULT 0, status TEXT NOT NULL DEFAULT 'POSTED', user_id INTEGER, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS purchase_items (
  id INTEGER PRIMARY KEY, purchase_id INTEGER NOT NULL REFERENCES purchases(id), product_id INTEGER NOT NULL REFERENCES products(id),
  qty_input REAL NOT NULL, sell_unit TEXT NOT NULL, qty_base REAL NOT NULL, cost_price REAL NOT NULL, gst_rate REAL NOT NULL,
  gst_amount REAL NOT NULL, batch_number TEXT, batch_id INTEGER REFERENCES stock_batches(id)
);
CREATE TABLE IF NOT EXISTS customers (
  id INTEGER PRIMARY KEY, customer_code TEXT UNIQUE NOT NULL, name TEXT NOT NULL, phone TEXT, address TEXT, gstin TEXT,
  customer_type TEXT NOT NULL DEFAULT 'RETAIL', credit_enabled INTEGER NOT NULL DEFAULT 0, credit_limit REAL NOT NULL DEFAULT 0,
  state_code TEXT, notes TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_customer_phone ON customers(phone);
CREATE TABLE IF NOT EXISTS invoices (
  id INTEGER PRIMARY KEY, invoice_number TEXT UNIQUE, invoice_date TEXT NOT NULL, customer_id INTEGER REFERENCES customers(id),
  customer_name_snapshot TEXT NOT NULL, customer_type_snapshot TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'DRAFT',
  pricing_mode TEXT NOT NULL DEFAULT 'RETAIL', subtotal REAL NOT NULL DEFAULT 0, discount REAL NOT NULL DEFAULT 0,
  taxable_amount REAL NOT NULL DEFAULT 0, cgst REAL NOT NULL DEFAULT 0, sgst REAL NOT NULL DEFAULT 0, igst REAL NOT NULL DEFAULT 0,
  cess REAL NOT NULL DEFAULT 0, grand_total REAL NOT NULL DEFAULT 0, amount_received REAL NOT NULL DEFAULT 0,
  outstanding REAL NOT NULL DEFAULT 0, user_id INTEGER, notes TEXT, created_at TEXT NOT NULL, finalized_at TEXT
);
CREATE INDEX IF NOT EXISTS idx_invoice_date_status ON invoices(invoice_date, status);
CREATE TABLE IF NOT EXISTS invoice_items (
  id INTEGER PRIMARY KEY, invoice_id INTEGER NOT NULL REFERENCES invoices(id), product_id INTEGER NOT NULL REFERENCES products(id),
  sku_snapshot TEXT NOT NULL, product_name_snapshot TEXT NOT NULL, hsn_snapshot TEXT, qty_input REAL NOT NULL, sell_unit TEXT NOT NULL,
  qty_base REAL NOT NULL, unit_price REAL NOT NULL, original_price REAL NOT NULL, discount REAL NOT NULL,
  taxable_amount REAL NOT NULL, gst_rate REAL NOT NULL, cess_rate REAL NOT NULL, gst_amount REAL NOT NULL,
  cost_snapshot REAL NOT NULL, manual_override_reason TEXT
);
CREATE TABLE IF NOT EXISTS payments (
  id INTEGER PRIMARY KEY, invoice_id INTEGER REFERENCES invoices(id), customer_id INTEGER REFERENCES customers(id),
  payment_date TEXT NOT NULL, mode TEXT NOT NULL, amount REAL NOT NULL, reference TEXT, notes TEXT, user_id INTEGER, created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_payments_invoice ON payments(invoice_id);
CREATE TABLE IF NOT EXISTS returns (
  id INTEGER PRIMARY KEY, return_number TEXT UNIQUE NOT NULL, invoice_id INTEGER NOT NULL REFERENCES invoices(id),
  customer_id INTEGER REFERENCES customers(id), return_date TEXT NOT NULL, reason TEXT, total_refund REAL NOT NULL DEFAULT 0,
  status TEXT NOT NULL DEFAULT 'POSTED', user_id INTEGER, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS return_items (
  id INTEGER PRIMARY KEY, return_id INTEGER NOT NULL REFERENCES returns(id), invoice_item_id INTEGER NOT NULL REFERENCES invoice_items(id),
  product_id INTEGER NOT NULL REFERENCES products(id), qty_input REAL NOT NULL, qty_base REAL NOT NULL, refund_amount REAL NOT NULL
);
CREATE TABLE IF NOT EXISTS expenses (
  id INTEGER PRIMARY KEY, expense_date TEXT NOT NULL, category TEXT NOT NULL, amount REAL NOT NULL, payment_mode TEXT NOT NULL,
  vendor TEXT, notes TEXT, attachment_name TEXT, recurring INTEGER NOT NULL DEFAULT 0, user_id INTEGER, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS recurring_expenses (
  id INTEGER PRIMARY KEY, category TEXT NOT NULL, amount REAL NOT NULL, payment_mode TEXT NOT NULL, vendor TEXT,
  frequency TEXT NOT NULL, start_date TEXT NOT NULL, end_date TEXT, active INTEGER NOT NULL DEFAULT 1, notes TEXT, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS held_bills (
  id INTEGER PRIMARY KEY, ref TEXT UNIQUE NOT NULL, customer_id INTEGER REFERENCES customers(id), customer_name TEXT,
  pricing_mode TEXT NOT NULL, items_json TEXT NOT NULL, notes TEXT, created_at TEXT NOT NULL, user_id INTEGER
);
CREATE TABLE IF NOT EXISTS audit_logs (
  id INTEGER PRIMARY KEY, user_id INTEGER, action TEXT NOT NULL, entity_type TEXT, entity_id INTEGER,
  previous_json TEXT, new_json TEXT, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS targets (
  id INTEGER PRIMARY KEY CHECK(id=1), daily REAL DEFAULT 500000, weekly REAL DEFAULT 3000000,
  monthly REAL DEFAULT 10000000, seasonal REAL DEFAULT 30000000, season_start TEXT, season_end TEXT
);
CREATE TABLE IF NOT EXISTS sequence (
  key TEXT PRIMARY KEY, value INTEGER NOT NULL
);
'''


def init_db():
    conn = db(); exec_script(conn, SCHEMA)
    product_columns={r['name'] for r in conn.execute('PRAGMA table_info(products)').fetchall()}
    user_columns={r['name'] for r in conn.execute('PRAGMA table_info(users)').fetchall()}
    if 'session_version' not in user_columns:
        conn.execute('ALTER TABLE users ADD COLUMN session_version INTEGER NOT NULL DEFAULT 0')
    invoice_columns={r['name'] for r in conn.execute('PRAGMA table_info(invoices)').fetchall()}
    for column, definition in {
        'revised_from_id':'INTEGER', 'revised_to_id':'INTEGER', 'revision_reason':'TEXT',
        'edited_by':'INTEGER', 'edited_at':'TEXT', 'voided_by':'INTEGER', 'voided_at':'TEXT'
    }.items():
        if column not in invoice_columns:
            conn.execute(f'ALTER TABLE invoices ADD COLUMN {column} {definition}')
    for column, definition in {
        'cost_price':'REAL NOT NULL DEFAULT 0', 'mrp':'REAL NOT NULL DEFAULT 0',
        'retail_selling_price':'REAL NOT NULL DEFAULT 0', 'wholesale_selling_price':'REAL NOT NULL DEFAULT 0',
        'default_unit_price':'REAL NOT NULL DEFAULT 0'
    }.items():
        if column not in product_columns:
            conn.execute(f'ALTER TABLE products ADD COLUMN {column} {definition}')
    conn.execute('''
        UPDATE products SET cost_price=COALESCE((SELECT cost_price FROM price_versions pv WHERE pv.product_id=products.id ORDER BY effective_date DESC,id DESC LIMIT 1),0),
          mrp=COALESCE((SELECT mrp FROM price_versions pv WHERE pv.product_id=products.id ORDER BY effective_date DESC,id DESC LIMIT 1),0),
          retail_selling_price=COALESCE((SELECT retail_price FROM price_versions pv WHERE pv.product_id=products.id ORDER BY effective_date DESC,id DESC LIMIT 1),0),
                    wholesale_selling_price=COALESCE((SELECT wholesale_price FROM price_versions pv WHERE pv.product_id=products.id ORDER BY effective_date DESC,id DESC LIMIT 1),0),
                    default_unit_price=CASE WHEN default_unit_price=0 THEN COALESCE((SELECT retail_price FROM price_versions pv WHERE pv.product_id=products.id ORDER BY effective_date DESC,id DESC LIMIT 1),0) ELSE default_unit_price END
                WHERE cost_price=0 AND mrp=0 AND retail_selling_price=0 AND wholesale_selling_price=0 AND default_unit_price=0
    ''')
    conn.commit()
    cur = conn.execute('SELECT COUNT(*) c FROM users')
    if cur.fetchone()['c'] == 0:
        admin_password=os.getenv('FIREWORKS_ADMIN_PASSWORD')
        if os.getenv('FIREWORKS_REQUIRE_ADMIN_PASSWORD') == '1' and not admin_password:
            conn.close(); raise RuntimeError('Set FIREWORKS_ADMIN_PASSWORD before initializing the hosted database.')
        if admin_password and len(admin_password)<12:
            conn.close(); raise RuntimeError('FIREWORKS_ADMIN_PASSWORD must be at least 12 characters.')
        if not admin_password and not is_loopback_host(HOST):
            conn.close(); raise RuntimeError('Set FIREWORKS_ADMIN_PASSWORD before exposing the app on a network (HOST is not localhost).')
        conn.execute('INSERT INTO users(username,password_hash,role,created_at) VALUES(?,?,?,?)', ('admin', hash_password(admin_password or 'admin123'), 'OWNER', now_iso()))
    if conn.execute('SELECT COUNT(*) c FROM business_settings').fetchone()['c'] == 0:
        settings = {
            'business_name':'Shree Ganesh Fireworks', 'address':'Pune, Maharashtra, India', 'gstin':'', 'license_number':'',
            'phone':'', 'email':'', 'logo':'', 'state':'Maharashtra', 'state_code':'27', 'invoice_prefix':'FW',
            'negative_stock_allowed':False, 'batch_tracking':True, 'default_gst_rate':0, 'default_cess_rate':0,
            'invoice_footer':'Thank you for shopping with us.', 'walk_in_name':'Walk-in Customer'
        }
        conn.execute('INSERT INTO business_settings(id,data_json) VALUES(1,?)', (json.dumps(settings),))
    if conn.execute('SELECT COUNT(*) c FROM targets').fetchone()['c'] == 0:
        conn.execute('INSERT INTO targets(id,daily,weekly,monthly,seasonal,season_start,season_end) VALUES(1,500000,3000000,10000000,30000000,?,?)',(f'{dt.date.today().year}-10-01',f'{dt.date.today().year}-10-31'))
    if conn.execute('SELECT COUNT(*) c FROM sequence').fetchone()['c'] == 0:
        conn.execute('INSERT INTO sequence(key,value) VALUES("invoice",0),("purchase",0),("return",0)')
    conn.commit(); conn.close()
    if os.getenv('FIREWORKS_SEED_DEMO') == '1':
        seed_demo()


def seed_demo():
    conn=db()
    if conn.execute('SELECT COUNT(*) c FROM products').fetchone()['c'] > 0:
        conn.close(); return
    cats = [('Sparklers','Ground'),('Aerial Shots','Aerial'),('Rockets','Aerial'),('Gift Boxes','Combo'),('Ground Spinners','Ground'),('Flower Pots','Ground'),('Sound Crackers','Crackers')]
    names=[]
    for cat, sub in cats:
        for i in range(5): names.append((cat, sub, f'{cat[:-1] if cat.endswith("s") else cat} {i+1}'))
    user_id=conn.execute('SELECT id FROM users WHERE username="admin"').fetchone()['id']
    for idx,(cat,sub,nm) in enumerate(names,1):
        sku=f'FW-{idx:04d}'
        unit='Box'
        upb=10 if cat in ('Sparklers','Ground Spinners') else 1
        bpc=12 if cat in ('Gift Boxes','Aerial Shots','Rockets') else 10
        upc=upb*bpc
        cost=round(40+idx*7.25,2); retail=round(cost*2.15,2); wholesale=round(cost*1.85,2); mrp=round(retail*1.2,2)
        conn.execute('INSERT INTO products(sku,name,category,subcategory,unit,units_per_box,boxes_per_carton,units_per_carton,brand,hsn,gst_rate,cess_rate,reorder_level,status,cost_price,mrp,retail_selling_price,wholesale_selling_price,default_unit_price,barcode,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
            (sku,nm,cat,sub,unit,upb,bpc,upc,'Demo Brand','3604',18,0,20,'ACTIVE',cost,mrp,retail,wholesale,retail,f'8900000{idx:04d}',now_iso(),now_iso()))
        pid=conn.execute('SELECT last_insert_rowid()').fetchone()[0]
        conn.execute('INSERT INTO price_versions(product_id,effective_date,mrp,retail_price,wholesale_price,cost_price,created_by,created_at) VALUES(?,?,?,?,?,?,?,?)', (pid, today(), mrp, retail, wholesale, cost, user_id, now_iso()))
        opening=25+idx*2
        base=opening*upb
        conn.execute('INSERT INTO inventory_movements(product_id,movement_type,qty_base,unit_cost,ref_type,reason,user_id,created_at) VALUES(?,?,?,?,?,?,?,?)', (pid,'OPENING',base,cost,'OPENING','Demo opening stock',user_id,now_iso()))
    for i,(name,ctype,phone,credit) in enumerate([
        ('Aarav Traders','WHOLESALE','9876500011',500000),('Mahalaxmi Distributors','WHOLESALE','9876500012',300000),
        ('Rohan Patil','RETAIL','9876500013',0),('Sneha Kulkarni','RETAIL','9876500014',0),('Walk-in Customer','RETAIL','',0)],1):
        conn.execute('INSERT INTO customers(customer_code,name,phone,customer_type,credit_enabled,credit_limit,state_code,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)', (f'CUST-{i:04d}',name,phone,ctype,1 if credit else 0,credit,'27',now_iso(),now_iso()))
    # Historical/demo invoices over the last 21 days, with heavier activity today.
    products=conn.execute('SELECT p.id,p.sku,p.name,p.category,p.unit,p.units_per_box,p.gst_rate,p.hsn, pv.retail_price,pv.wholesale_price,pv.cost_price,pv.mrp FROM products p JOIN price_versions pv ON pv.product_id=p.id AND pv.id=(SELECT id FROM price_versions x WHERE x.product_id=p.id ORDER BY effective_date DESC,id DESC LIMIT 1) ORDER BY p.id').fetchall()
    customers=conn.execute('SELECT * FROM customers').fetchall()
    for d_off in range(20,-1,-1):
        d=(dt.date.today()-dt.timedelta(days=d_off)).isoformat()
        bills=5 if d_off>2 else (8 if d_off>0 else 24)
        for b in range(bills):
            cust=customers[(b+d_off)%len(customers)]
            mode='WHOLESALE' if cust['customer_type']=='WHOLESALE' else 'RETAIL'
            subtotal=0; disc=0; taxable=0; cgst=sgst=igst=cess=0; cogs=0
            rows=[]
            for j in range(1 + (b%3)):
                p=products[(b*3+j+d_off)%len(products)]
                qty=1+(b+j)%4
                price=p['wholesale_price'] if mode=='WHOLESALE' else p['retail_price']
                od=round(price*qty,2); dd=round(od*(0.02 if b%5==0 else 0),2); tx=od-dd
                tax=round(tx*p['gst_rate']/100,2); c1=round(tx*0.5*p['gst_rate']/100,2); c2=tax-c1
                subtotal+=od; disc+=dd; taxable+=tx; cgst+=c1; sgst+=c2; cogs+=p['cost_price']*qty
                rows.append((p,qty,price,dd,tx,tax,cogs))
            total=round(taxable+cgst+sgst+igst+cess,2)
            paid=total if (b+d_off)%4!=0 else round(total*0.35,2)
            invoice_num=next_seq(conn,'invoice',d)
            conn.execute('INSERT INTO invoices(invoice_number,invoice_date,customer_id,customer_name_snapshot,customer_type_snapshot,status,pricing_mode,subtotal,discount,taxable_amount,cgst,sgst,igst,cess,grand_total,amount_received,outstanding,user_id,created_at,finalized_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                (invoice_num,d,cust['id'],cust['name'],cust['customer_type'],'FINALIZED',mode,subtotal,disc,taxable,cgst,sgst,igst,cess,total,paid,round(total-paid,2),user_id,f'{d}T{10+(b%10):02d}:{(b*7)%60:02d}:00',f'{d}T{10+(b%10):02d}:{(b*7)%60:02d}:00'))
            iid=conn.execute('SELECT last_insert_rowid()').fetchone()[0]
            for p,qty,price,dd,tx,tax,_ in rows:
                conn.execute('INSERT INTO invoice_items(invoice_id,product_id,sku_snapshot,product_name_snapshot,hsn_snapshot,qty_input,sell_unit,qty_base,unit_price,original_price,discount,taxable_amount,gst_rate,cess_rate,gst_amount,cost_snapshot) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                    (iid,p['id'],p['sku'],p['name'],p['hsn'],qty,p['unit'],qty*p['units_per_box'],price,price,dd,tx,p['gst_rate'],0,tax,p['cost_price']))
                conn.execute('INSERT INTO inventory_movements(product_id,movement_type,qty_base,unit_cost,ref_type,ref_id,user_id,created_at) VALUES(?,?,?,?,?,?,?,?)',
                    (p['id'],'SALE',-qty*p['units_per_box'],p['cost_price'],'INVOICE',iid,user_id,f'{d}T{10+(b%10):02d}:{(b*7)%60:02d}:00'))
            if paid:
                conn.execute('INSERT INTO payments(invoice_id,customer_id,payment_date,mode,amount,reference,user_id,created_at) VALUES(?,?,?,?,?,?,?,?)', (iid,cust['id'],d,'UPI' if b%2 else 'Cash',paid,f'DEMO-{iid}',user_id,f'{d}T{10+(b%10):02d}:{(b*7)%60:02d}:00'))
    # A few expenses
    for off,cat,amt in [(0,'Transport/Freight',18000),(1,'Salaries/Wages',12000),(2,'Packaging',4500),(3,'Marketing',7500),(7,'Rent',25000)]:
        conn.execute('INSERT INTO expenses(expense_date,category,amount,payment_mode,vendor,notes,user_id,created_at) VALUES(?,?,?,?,?,?,?,?)', ((dt.date.today()-dt.timedelta(days=off)).isoformat(),cat,amt,'UPI','Demo Vendor','Sample/Test data',user_id,now_iso()))
    # one return
    inv=conn.execute("SELECT * FROM invoices ORDER BY id DESC LIMIT 1 OFFSET 2").fetchone()
    if inv:
        item=conn.execute('SELECT * FROM invoice_items WHERE invoice_id=? LIMIT 1',(inv['id'],)).fetchone()
        if item:
            ret_num=next_seq(conn,'return',today())
            conn.execute('INSERT INTO returns(return_number,invoice_id,customer_id,return_date,reason,total_refund,user_id,created_at) VALUES(?,?,?,?,?,?,?,?)',(ret_num,inv['id'],inv['customer_id'],today(),'Demo partial return',item['unit_price'],user_id,now_iso()))
            rid=conn.execute('SELECT last_insert_rowid()').fetchone()[0]
            conn.execute('INSERT INTO return_items(return_id,invoice_item_id,product_id,qty_input,qty_base,refund_amount) VALUES(?,?,?,?,?,?)',(rid,item['id'],item['product_id'],1,item['qty_base']/item['qty_input'],item['unit_price']))
            conn.execute('INSERT INTO inventory_movements(product_id,movement_type,qty_base,unit_cost,ref_type,ref_id,reason,user_id,created_at) VALUES(?,?,?,?,?,?,?,?,?)',(item['product_id'],'RETURN_IN',item['qty_base']/item['qty_input'],item['cost_snapshot'],'RETURN',rid,'Demo return',user_id,now_iso()))
    conn.commit(); conn.close()


def next_seq(conn, key, day=None):
    row=conn.execute('SELECT value FROM sequence WHERE key=?',(key,)).fetchone(); val=(row['value'] if row else 0)+1
    conn.execute('INSERT INTO sequence(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value',(key,val))
    prefix={'invoice':'FW','purchase':'PUR','return':'RET'}[key]
    return f'{prefix}-{(day or today()).replace("-","")}-{val:05d}'


def settings(conn):
    return json.loads(conn.execute('SELECT data_json FROM business_settings WHERE id=1').fetchone()['data_json'])


def get_price(conn,pid,effective=None):
    q='SELECT * FROM price_versions WHERE product_id=? AND effective_date<=? ORDER BY effective_date DESC,id DESC LIMIT 1'
    row=conn.execute(q,(pid,effective or today())).fetchone()
    return dict(row) if row else None


def product_dict(conn, p):
    d=dict(p)
    for legacy in ('units_per_box','boxes_per_carton','units_per_carton','hsn','gst_rate','cess_rate','cost_price','mrp','retail_selling_price','wholesale_selling_price','barcode'):
        d.pop(legacy,None)
    stock=conn.execute('SELECT COALESCE(SUM(qty_base),0) q FROM inventory_movements WHERE product_id=?',(p['id'],)).fetchone()['q']
    d['current_stock_base']=stock
    d['default_unit_price']=p['default_unit_price']
    d['price']={'unit_price':p['default_unit_price']}
    return d


def audit(conn,user_id,action,etype,eid,old=None,new=None):
    conn.execute('INSERT INTO audit_logs(user_id,action,entity_type,entity_id,previous_json,new_json,created_at) VALUES(?,?,?,?,?,?,?)', (user_id,action,etype,eid,json.dumps(old,default=str) if old else None,json.dumps(new,default=str) if new else None,now_iso()))


def read_body(handler):
    length=int(handler.headers.get('Content-Length','0') or 0)
    if length<0 or length>MAX_BODY_BYTES:
        raise ValueError(f'Request is too large (limit {MAX_BODY_BYTES//(1024*1024)} MB).')
    return handler.rfile.read(length) if length else b''


def parse_body(handler):
    raw=read_body(handler)
    ctype=handler.headers.get('Content-Type','')
    if 'application/json' in ctype:
        return json.loads(raw or b'{}')
    if 'application/x-www-form-urlencoded' in ctype:
        return {k:v[0] for k,v in urllib.parse.parse_qs(raw.decode()).items()}
    return raw


def parse_multipart(handler):
    raw=read_body(handler)
    ctype=handler.headers.get('Content-Type','')
    header=b'Content-Type: '+ctype.encode()+b'\r\nMIME-Version: 1.0\r\n\r\n'
    msg=BytesParser(policy=default).parsebytes(header+raw)
    fields={}; files={}
    for part in msg.iter_parts():
        name=part.get_param('name',header='content-disposition')
        fn=part.get_filename()
        payload=part.get_payload(decode=True) or b''
        if fn: files[name]={'filename':fn,'content':payload}
        else: fields[name]=payload.decode('utf-8',errors='replace')
    return fields,files


def csv_or_xlsx_rows(file_info):
    data=file_info['content']; fn=file_info['filename'].lower()
    if fn.endswith('.csv'):
        text=data.decode('utf-8-sig',errors='replace'); return list(csv.DictReader(io.StringIO(text)))
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    ws=wb.active; rows=list(ws.iter_rows(values_only=True)); wb.close()
    if not rows: return []
    headers=[str(x).strip() if x is not None else '' for x in rows[0]]
    out=[]
    for r in rows[1:]: out.append({headers[i]: (r[i] if i<len(r) else None) for i in range(len(headers))})
    return out

COLUMN_ALIASES={
 'sku':['sku','product code','product_code','productcode','product id','product_id','serial number','serial no','serial no.','product number','product no'],
 'name':['product name','name','product'], 'category':['category'], 'subcategory':['subcategory','sub-category'],
 'brand':['brand / manufacturer','brand','manufacturer'], 'unit':['unit','uom','selling unit'],
 'unit_price':['unit price','unit selling price','selling price','default unit price','default selling price','retail selling price','retail price','price'],
 'opening_stock':['opening stock','opening_stock','stock','stock quantity','opening quantity'],
 'reorder_level':['minimum stock / reorder level','reorder level','reorder_level','minimum stock'],
 'status':['product status','status']
}

def norm(s): return re.sub(r'[^a-z0-9]+',' ',str(s or '').strip().lower()).strip()

def suggest_map(headers):
    nheaders={norm(h):h for h in headers}; mapping={}
    for key,aliases in COLUMN_ALIASES.items():
        best=None; best_score=0
        for a in aliases:
            na=norm(a)
            for nh,h in nheaders.items():
                if nh==na: score=100
                elif na in nh or nh in na: score=80 + min(len(na),len(nh))/100
                else:
                    at=set(na.split()); ht=set(nh.split()); overlap=len(at & ht)
                    score=overlap/(max(1,len(at))) * 70
                if score>best_score:
                    best_score=score; best=h
        mapping[key]=best if best_score>=45 else None
    return mapping


def num(v, default=0):
    if v is None or v=='': return default
    s=str(v).replace(',','').replace('₹','').strip()
    try: return float(s)
    except: return default


def stock_base_factor(p, sell_unit):
    return 1


def price_for_unit(p, price, sell_unit):
    return price


def serialize_rows(rows): return [dict(r) for r in rows]


def current_stock(conn,pid): return conn.execute('SELECT COALESCE(SUM(qty_base),0) q FROM inventory_movements WHERE product_id=?',(pid,)).fetchone()['q']


def encode_held_bill(payload):
    return json.dumps({
        'items': payload.get('items') or [],
        'discount': money(payload.get('discount',0)),
        'customer_phone': str(payload.get('customer_phone') or '').strip()
    })


def decode_held_bill(value):
    stored=json.loads(value or '[]')
    if isinstance(stored,list):
        return {'items':stored,'discount':0.0,'customer_phone':''}
    if isinstance(stored,dict):
        return {'items':stored.get('items') or [],'discount':money(stored.get('discount',0)),'customer_phone':str(stored.get('customer_phone') or '')}
    return {'items':[],'discount':0.0,'customer_phone':''}


def consume_held_bill(conn, held_bill_id):
    row=conn.execute('SELECT * FROM held_bills WHERE id=?',(held_bill_id,)).fetchone()
    if not row: raise ValueError('Held bill not found.')
    result={'bill':dict(row),**decode_held_bill(row['items_json'])}
    conn.execute('DELETE FROM held_bills WHERE id=?',(held_bill_id,))
    return result


def create_invoice(conn, payload, user_id, held=False, invoice_date=None, billed_by=None):
    customer_id=payload.get('customer_id'); mode='STANDARD'
    cust=conn.execute('SELECT * FROM customers WHERE id=?',(customer_id,)).fetchone() if customer_id else None
    manual_name=str(payload.get('customer_name') or '').strip()
    manual_phone=str(payload.get('customer_phone') or '').strip()
    if manual_phone and not re.fullmatch(r'\d{10}', manual_phone):
        raise ValueError('Mobile number must be exactly 10 digits.')
    if not cust and (manual_name or manual_phone):
        if manual_phone:
            cust=conn.execute('SELECT * FROM customers WHERE phone=? ORDER BY id LIMIT 1',(manual_phone,)).fetchone()
        if not cust and manual_name:
            cust=conn.execute('SELECT * FROM customers WHERE name=? ORDER BY id LIMIT 1',(manual_name,)).fetchone()
        if not cust:
            code=f'CUST-{secrets.randbelow(99999):05d}'
            conn.execute('INSERT INTO customers(customer_code,name,phone,customer_type,credit_enabled,credit_limit,state_code,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)',(code,manual_name or 'Customer',manual_phone,'CUSTOMER',0,0,settings(conn).get('state_code'),now_iso(),now_iso()))
            customer_id=conn.execute('SELECT last_insert_rowid()').fetchone()[0]
            cust=conn.execute('SELECT * FROM customers WHERE id=?',(customer_id,)).fetchone()
        else:
            customer_id=cust['id']
    cust_name=(cust['name'] if cust else (payload.get('customer_name') or settings(conn)['walk_in_name'])); cust_type=(cust['customer_type'] if cust else 'CUSTOMER')
    items=payload.get('items') or []
    if not items: raise ValueError('Add at least one product to the bill.')
    s=settings(conn); rows=[]; subtotal=discount=taxable=cgst=sgst=igst=cess=cogs=0
    for it in items:
        pid=int(it['product_id']); p=conn.execute('SELECT * FROM products WHERE id=? AND status="ACTIVE"',(pid,)).fetchone()
        if not p: raise ValueError(f'Product {pid} not found or inactive.')
        sell_unit=p['unit']; qty=float(it.get('qty',0))
        if qty<=0: raise ValueError(f'Invalid quantity for {p["name"]}.')
        qty_base=qty
        stock=current_stock(conn,pid)
        if stock < qty_base and not s.get('negative_stock_allowed',False): raise ValueError(f'Insufficient stock for {p["name"]}. Available: {stock:g} base units.')
        original=float(p['default_unit_price'] or 0)
        if original<0: raise ValueError(f'Invalid default unit price for {p["name"]}.')
        final=original
        if final<0: raise ValueError('Price cannot be negative.')
        line_gross=final*qty
        line_disc=0; line_total=line_gross
        subtotal+=line_total; taxable+=line_total
        rows.append((p,qty,sell_unit,qty_base,final,original,line_disc,line_total,0,0,0,0,0,0,0,0))
    subtotal=money(subtotal)
    try:
        discount=money(payload.get('discount',0))
    except (ValueError, TypeError):
        raise ValueError('Discount must be a valid amount.')
    if not math.isfinite(discount) or discount < 0 or discount > subtotal:
        raise ValueError('Discount must be between 0 and the subtotal.')
    line_discounts=allocate_discount(discount,[row[7] for row in rows])
    rows=[tuple(row[:6]+(line_discounts[index],money(row[7]-line_discounts[index]))+row[8:]) for index,row in enumerate(rows)]
    taxable=money(sum(row[7] for row in rows))
    grand=taxable
    payments=payload.get('payments') or []
    total_pay=credit_total=0
    for payment in payments:
        amount=float(payment.get('amount',0) or 0)
        if not math.isfinite(amount) or amount<0:
            raise ValueError('Payment amounts must be valid non-negative amounts.')
        total_pay+=amount
        if str(payment.get('mode','Cash')).strip().upper()=='CREDIT':
            credit_total+=amount
    if abs(total_pay-grand)>0.01:
        raise ValueError('Payment allocations must equal the invoice total after discount.')
    received=money(total_pay-credit_total)
    outstanding=money(credit_total)
    inv_num=next_seq(conn,'invoice')
    conn.execute('INSERT INTO invoices(invoice_number,invoice_date,customer_id,customer_name_snapshot,customer_type_snapshot,status,pricing_mode,subtotal,discount,taxable_amount,cgst,sgst,igst,cess,grand_total,amount_received,outstanding,user_id,notes,created_at,finalized_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                 (inv_num,invoice_date or today(),customer_id,cust_name,cust_type,'FINALIZED',mode,subtotal,discount,taxable,cgst,sgst,igst,cess,grand,received,outstanding,billed_by or user_id,payload.get('notes'),now_iso(),now_iso()))
    iid=conn.execute('SELECT last_insert_rowid()').fetchone()[0]
    for p,qty,su,qb,final,original,ld,lt,tax,ca,cost,lcg,lsg,lig,gr,cr in rows:
        conn.execute('INSERT INTO invoice_items(invoice_id,product_id,sku_snapshot,product_name_snapshot,hsn_snapshot,qty_input,sell_unit,qty_base,unit_price,original_price,discount,taxable_amount,gst_rate,cess_rate,gst_amount,cost_snapshot,manual_override_reason) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                 (iid,p['id'],p['sku'],p['name'],p['hsn'],qty,su,qb,final,original,ld,lt,0,0,0,cost,payload.get('override_reason')))
        conn.execute('INSERT INTO inventory_movements(product_id,movement_type,qty_base,unit_cost,ref_type,ref_id,reason,user_id,created_at) VALUES(?,?,?,?,?,?,?,?,?)',
                     (p['id'],'SALE',-qb,cost,'INVOICE',iid,'Sale',user_id,now_iso()))
    for pay in payments:
        amt=float(pay.get('amount',0) or 0)
        if amt:
            conn.execute('INSERT INTO payments(invoice_id,customer_id,payment_date,mode,amount,reference,notes,user_id,created_at) VALUES(?,?,?,?,?,?,?,?,?)', (iid,customer_id,today(),pay.get('mode','Cash'),amt,pay.get('reference'),pay.get('notes'),user_id,now_iso()))
    audit(conn,user_id,'Invoice finalized','INVOICE',iid,new={'invoice_number':inv_num,'total':grand})
    return iid


def void_invoice(conn, invoice_id, user_id):
    inv=conn.execute('SELECT * FROM invoices WHERE id=?',(invoice_id,)).fetchone()
    if not inv or inv['status']!='FINALIZED': raise ValueError('Finalized invoice not found or already voided.')
    returns=conn.execute('SELECT COUNT(*) c FROM returns WHERE invoice_id=? AND status="POSTED"',(invoice_id,)).fetchone()['c']
    if returns: raise ValueError('This sale has a posted return and cannot be deleted.')
    items=conn.execute('SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY id',(invoice_id,)).fetchall()
    for item in items:
        conn.execute('INSERT INTO inventory_movements(product_id,movement_type,qty_base,unit_cost,ref_type,ref_id,reason,user_id,created_at) VALUES(?,?,?,?,?,?,?,?,?)',
                     (item['product_id'],'SALE_VOID',item['qty_base'],item['cost_snapshot'],'INVOICE',invoice_id,'Voided sale stock restoration',user_id,now_iso()))
    conn.execute('UPDATE invoices SET status="VOID",voided_by=?,voided_at=? WHERE id=?',(user_id,now_iso(),invoice_id))
    audit(conn,user_id,'Invoice voided','INVOICE',invoice_id,old=dict(inv),new={'status':'VOID','stock_restored':len(items)})
    return {'invoice_id':invoice_id,'restored_items':len(items)}


def edit_invoice(conn, invoice_id, payload, user_id):
    reason=str(payload.get('reason') or '').strip()
    if len(reason)<3: raise ValueError('Enter a reason for editing this bill.')
    old=conn.execute('SELECT * FROM invoices WHERE id=?',(invoice_id,)).fetchone()
    if not old or old['status']!='FINALIZED': raise ValueError('Only a finalized bill can be edited.')
    old_items=[dict(r) for r in conn.execute('SELECT product_id,product_name_snapshot,qty_input FROM invoice_items WHERE invoice_id=? ORDER BY id',(invoice_id,)).fetchall()]
    void_invoice(conn,invoice_id,user_id)
    new_id=create_invoice(conn,payload,user_id,invoice_date=old['invoice_date'],billed_by=old['user_id'])
    edited_at=now_iso()
    conn.execute('UPDATE invoices SET revised_from_id=?,revision_reason=?,edited_by=?,edited_at=? WHERE id=?',(invoice_id,reason,user_id,edited_at,new_id))
    conn.execute('UPDATE invoices SET revised_to_id=? WHERE id=?',(new_id,invoice_id))
    new=conn.execute('SELECT invoice_number,grand_total FROM invoices WHERE id=?',(new_id,)).fetchone()
    new_items=[dict(r) for r in conn.execute('SELECT product_id,product_name_snapshot,qty_input FROM invoice_items WHERE invoice_id=? ORDER BY id',(new_id,)).fetchall()]
    audit(conn,user_id,'Invoice edited','INVOICE',new_id,
        old={'invoice_id':invoice_id,'invoice_number':old['invoice_number'],'total':old['grand_total'],'items':old_items},
        new={'invoice_id':new_id,'invoice_number':new['invoice_number'],'total':new['grand_total'],'items':new_items,'reason':reason})
    return {'invoice_id':new_id,'invoice_number':new['invoice_number'],'previous_invoice_number':old['invoice_number'],'edited_at':edited_at}


INVOICE_SELECT='''SELECT i.*, ub.username billed_by_name, ue.username edited_by_name, uv.username voided_by_name,
  rf.invoice_number revised_from_number, rt.invoice_number revised_to_number
FROM invoices i
LEFT JOIN users ub ON ub.id=i.user_id LEFT JOIN users ue ON ue.id=i.edited_by LEFT JOIN users uv ON uv.id=i.voided_by
LEFT JOIN invoices rf ON rf.id=i.revised_from_id LEFT JOIN invoices rt ON rt.id=i.revised_to_id'''


def build_backup(conn):
    tables=[r['name'] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'").fetchall()]
    data={}
    for t in tables:
        rows=conn.execute(f'SELECT * FROM {t}').fetchall(); data[t]=[dict(r) for r in rows]
    # Backups are downloadable by staff, so never include credentials.
    for user in data.get('users',[]):
        user.pop('password_hash',None); user.pop('session_version',None)
    return data

class Handler(BaseHTTPRequestHandler):
    server_version='FireworksCRM/1.0'
    def log_message(self, fmt, *args): return

    def send(self,status,ctype=None,data=None,extra_headers=None):
        if data is None and isinstance(ctype,(bytes,bytearray)):
            data=ctype; ctype='application/json'
        self.send_response(status); self.send_header('Content-Type',ctype); self.send_header('Cache-Control','no-store')
        for k,v in SECURITY_HEADERS.items(): self.send_header(k,v)
        if extra_headers:
            for k,v in extra_headers.items(): self.send_header(k,v)
        self.send_header('Content-Length',str(len(data))); self.end_headers(); self.wfile.write(data)

    def current_user(self):
        tok=self.headers.get('Cookie','')
        m=re.search(r'(^|;\s*)'+re.escape(COOKIE)+r'=([^;]+)',tok)
        if not m: return None
        p=verify_token(urllib.parse.unquote(m.group(2)))
        if not p: return None
        conn=db(); u=conn.execute('SELECT * FROM users WHERE id=? AND active=1',(p.get('uid'),)).fetchone(); conn.close()
        if not u or p.get('sv',0)!=u['session_version']: return None
        return dict(u)

    def same_origin(self):
        # Reject cross-site POSTs (CSRF). Browsers send Origin on POST; non-browser clients may omit it.
        origin=self.headers.get('Origin')
        if not origin: return True
        return urllib.parse.urlparse(origin).netloc==self.headers.get('Host','')

    def require_auth(self):
        u=self.current_user()
        if not u:
            self.send(*json_response({'error':'Authentication required'},401),extra_headers={'Content-Type':'application/json'}); return None
        return u

    def do_GET(self):
        path=urllib.parse.urlparse(self.path).path; qs=urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)
        if path=='/':
            data=open(os.path.join(STATIC_DIR,'index.html'),'rb').read(); self.send(200,'text/html; charset=utf-8',data); return
        if path.startswith('/static/'):
            fn=os.path.basename(path); fp=os.path.join(STATIC_DIR,fn)
            if os.path.exists(fp):
                mt='text/plain'
                if fn.endswith('.js'): mt='application/javascript'
                elif fn.endswith('.css'): mt='text/css'
                elif fn.endswith('.webmanifest'): mt='application/manifest+json'
                self.send(200,mt,open(fp,'rb').read()); return
            self.send(404,'text/plain',b'Not found'); return
        if path=='/favicon.ico': self.send(204,'text/plain',b''); return
        if path=='/api/me':
            u=self.current_user(); self.send(*json_response({'authenticated':bool(u),'user':safe_user(u)})); return
        if path=='/api/login': self.send(*json_response({'error':'Use POST'},405)); return
        u=self.require_auth()
        if not u: return
        if not is_admin_user(u) and not staff_api_allowed('GET', path):
            self.send(*json_response({'error':'This account does not have access to that feature.'},403)); return
        conn=db()
        try:
            if path=='/api/users':
                if not is_admin_user(u):
                    self.send(*json_response({'error':'Administrator access required.'},403)); return
                users=conn.execute('SELECT id,username,role,active,created_at FROM users ORDER BY role,username').fetchall()
                self.send(*json_response(serialize_rows(users))); return
            if path=='/api/settings':
                st=settings(conn); tg=dict(conn.execute('SELECT * FROM targets WHERE id=1').fetchone()); self.send(*json_response({'settings':st,'targets':tg})); return
            if path.startswith('/api/products/') and path.count('/')==3:
                pid=int(path.rsplit('/',1)[1]); p=conn.execute('SELECT * FROM products WHERE id=?',(pid,)).fetchone()
                if not p: self.send(*json_response({'error':'Product not found'},404)); return
                prices=conn.execute('SELECT * FROM price_versions WHERE product_id=? ORDER BY effective_date DESC,id DESC',(pid,)).fetchall()
                stock=current_stock(conn,pid)
                last_sale=conn.execute('SELECT MAX(i.invoice_date) d FROM invoice_items ii JOIN invoices i ON i.id=ii.invoice_id WHERE ii.product_id=? AND i.status="FINALIZED"',(pid,)).fetchone()['d']
                last_purchase=conn.execute('SELECT MAX(p.purchase_date) d FROM purchase_items pi JOIN purchases p ON p.id=pi.purchase_id WHERE pi.product_id=?',(pid,)).fetchone()['d']
                self.send(*json_response({'product':product_dict(conn,p),'prices':serialize_rows(prices),'stock':stock,'last_sale':last_sale,'last_purchase':last_purchase})); return
            if path=='/api/products':
                q=(qs.get('q',[''])[0]).strip(); cat=qs.get('category',[''])[0]; status=qs.get('status',['ACTIVE'])[0]; limit=min(int(qs.get('limit',['200'])[0]),500)
                where=['p.status=?']; args=[status]
                if q: where.append('(p.name LIKE ? OR p.sku LIKE ? OR p.barcode LIKE ? OR p.category LIKE ?)'); args += [f'%{q}%']*4
                if cat: where.append('p.category=?'); args.append(cat)
                rows=conn.execute(f'SELECT p.* FROM products p WHERE {" AND ".join(where)} ORDER BY p.name LIMIT ?',args+[limit]).fetchall()
                out=[product_dict(conn,r) for r in rows]; self.send(*json_response(out)); return
            if path=='/api/categories':
                rows=conn.execute('SELECT DISTINCT category FROM products WHERE status="ACTIVE" AND category IS NOT NULL ORDER BY category').fetchall(); self.send(*json_response([r['category'] for r in rows])); return
            if path=='/api/customers':
                q=qs.get('q',[''])[0]; rows=conn.execute('SELECT * FROM customers WHERE name LIKE ? OR phone LIKE ? OR customer_code LIKE ? ORDER BY name LIMIT 100',(f'%{q}%',f'%{q}%',f'%{q}%')).fetchall(); self.send(*json_response(serialize_rows(rows))); return
            if path=='/api/invoices':
                start=qs.get('start',[today()])[0]; end=qs.get('end',[today()])[0]; status=qs.get('status',['FINALIZED'])[0]
                rows=conn.execute(INVOICE_SELECT+' WHERE i.invoice_date BETWEEN ? AND ? AND i.status=? ORDER BY i.id DESC LIMIT 500',(start,end,status)).fetchall(); self.send(*json_response(serialize_rows(rows))); return
            if path.startswith('/api/invoices/'):
                iid=int(path.rsplit('/',1)[1]); inv=conn.execute(INVOICE_SELECT.replace('SELECT i.*,','SELECT i.*, c.phone customer_phone,').replace('FROM invoices i','FROM invoices i LEFT JOIN customers c ON c.id=i.customer_id')+' WHERE i.id=?',(iid,)).fetchone()
                if not inv: self.send(*json_response({'error':'Invoice not found'},404)); return
                items=conn.execute('SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY id',(iid,)).fetchall(); pays=conn.execute('SELECT * FROM payments WHERE invoice_id=? ORDER BY id',(iid,)).fetchall(); self.send(*json_response({'invoice':dict(inv),'items':serialize_rows(items),'payments':serialize_rows(pays)})); return
            if path=='/api/dashboard':
                day=qs.get('date',[today()])[0]; yday=(dt.date.fromisoformat(day)-dt.timedelta(days=1)).isoformat(); lweek=(dt.date.fromisoformat(day)-dt.timedelta(days=7)).isoformat()
                cur=sales_kpis(conn,day,day); yd=sales_kpis(conn,yday,yday); lw=sales_kpis(conn,lweek,lweek)
                tg=dict(conn.execute('SELECT * FROM targets WHERE id=1').fetchone())
                exp=expense_total(conn,day,day)
                low=conn.execute('SELECT COUNT(*) c FROM products p WHERE p.status="ACTIVE" AND (SELECT COALESCE(SUM(qty_base),0) FROM inventory_movements m WHERE m.product_id=p.id) <= p.reorder_level').fetchone()['c']
                inv_cost=conn.execute('SELECT COALESCE(SUM((SELECT cost_price FROM price_versions pv WHERE pv.product_id=p.id ORDER BY effective_date DESC,id DESC LIMIT 1)* (SELECT COALESCE(SUM(qty_base),0) FROM inventory_movements m WHERE m.product_id=p.id)),0) v FROM products p WHERE p.status="ACTIVE"').fetchone()['v']
                attention=conn.execute('SELECT p.name,p.sku,ROUND((SELECT COALESCE(SUM(qty_base),0) FROM inventory_movements m WHERE m.product_id=p.id),2) stock,ROUND(p.reorder_level,2) reorder_level FROM products p WHERE p.status="ACTIVE" AND (SELECT COALESCE(SUM(qty_base),0) FROM inventory_movements m WHERE m.product_id=p.id) <= p.reorder_level ORDER BY stock ASC,p.name LIMIT 10').fetchall()
                self.send(*json_response({'today':cur,'yesterday':yd,'last_week':lw,'expenses':exp,'low_stock':low,'inventory_cost':inv_cost,'target':tg['daily'],'attention':serialize_rows(attention)})); return
            if path=='/api/analytics':
                start=qs.get('start',[today()])[0]; end=qs.get('end',[today()])[0]
                if qs.get('preset',[''])[0]: start,end=resolve_preset(qs.get('preset',[''])[0])
                kpi=sales_kpis(conn,start,end); ex=expense_total(conn,start,end); hourly=conn.execute('SELECT substr(finalized_at,12,2) hr, ROUND(SUM(grand_total),2) sales, COUNT(*) bills, ROUND(AVG(grand_total),2) avg_bill FROM invoices WHERE invoice_date BETWEEN ? AND ? AND status="FINALIZED" GROUP BY hr ORDER BY hr',(start,end)).fetchall()
                top=conn.execute('SELECT product_name_snapshot name, sku_snapshot sku, ROUND(SUM(qty_input),2) qty, ROUND(SUM(ii.taxable_amount),2) revenue, ROUND(SUM(ii.taxable_amount - ii.cost_snapshot*ii.qty_input),2) profit FROM invoice_items ii JOIN invoices i ON i.id=ii.invoice_id WHERE i.invoice_date BETWEEN ? AND ? AND i.status="FINALIZED" GROUP BY product_id ORDER BY revenue DESC LIMIT 20',(start,end)).fetchall()
                cat=conn.execute('SELECT COALESCE(p.category,"Uncategorised") category, ROUND(SUM(ii.taxable_amount),2) revenue, ROUND(SUM(ii.qty_input),2) qty, ROUND(SUM(ii.taxable_amount - ii.cost_snapshot*ii.qty_input),2) profit FROM invoice_items ii JOIN invoices i ON i.id=ii.invoice_id JOIN products p ON p.id=ii.product_id WHERE i.invoice_date BETWEEN ? AND ? AND i.status="FINALIZED" GROUP BY p.category ORDER BY revenue DESC',(start,end)).fetchall()
                payments=conn.execute('SELECT mode, ROUND(SUM(amount),2) amount FROM payments pa JOIN invoices i ON i.id=pa.invoice_id WHERE i.invoice_date BETWEEN ? AND ? AND i.status="FINALIZED" GROUP BY mode ORDER BY amount DESC',(start,end)).fetchall()
                customers=conn.execute('SELECT i.customer_type_snapshot customer_type, ROUND(SUM(i.grand_total),2) revenue, COUNT(*) bills, COUNT(DISTINCT i.customer_id) customers FROM invoices i WHERE i.invoice_date BETWEEN ? AND ? AND i.status="FINALIZED" GROUP BY i.customer_type_snapshot',(start,end)).fetchall()
                slow=conn.execute('SELECT p.name,p.sku,ROUND((SELECT COALESCE(SUM(qty_base),0) FROM inventory_movements m WHERE m.product_id=p.id),2) stock, (SELECT MAX(i.invoice_date) FROM invoice_items ii JOIN invoices i ON i.id=ii.invoice_id WHERE ii.product_id=p.id AND i.status="FINALIZED") last_sale, (SELECT COALESCE(SUM(ii.qty_input),0) FROM invoice_items ii JOIN invoices i ON i.id=ii.invoice_id WHERE ii.product_id=p.id AND i.invoice_date BETWEEN ? AND ? AND i.status="FINALIZED") qty_sold FROM products p WHERE p.status="ACTIVE" ORDER BY qty_sold ASC, stock DESC LIMIT 20',(start,end)).fetchall()
                dd=conn.execute('SELECT ROUND(SUM(discount),2) total_discount, ROUND(CASE WHEN SUM(subtotal)=0 THEN 0 ELSE SUM(discount)/SUM(subtotal)*100 END,2) avg_discount FROM invoices WHERE invoice_date BETWEEN ? AND ? AND status="FINALIZED"',(start,end)).fetchone()
                self.send(*json_response({'period':{'start':start,'end':end},'kpi':kpi,'expenses':ex,'net_profit':kpi['gross_profit']-ex,'hourly':serialize_rows(hourly),'top_products':serialize_rows(top),'categories':serialize_rows(cat),'payments':serialize_rows(payments),'customers':serialize_rows(customers),'slow_movers':serialize_rows(slow),'bottom_products':serialize_rows(slow[:10]),'discounts':dict(dd),'target':dict(conn.execute('SELECT * FROM targets WHERE id=1').fetchone())})); return
            if path=='/api/recurring-expenses':
                rows=conn.execute('SELECT * FROM recurring_expenses ORDER BY id DESC').fetchall(); self.send(*json_response(serialize_rows(rows))); return
            if path=='/api/held-bills':
                rows=conn.execute('SELECT * FROM held_bills ORDER BY created_at DESC').fetchall(); out=[]
                for r in rows:
                    d=dict(r); held=decode_held_bill(d.pop('items_json')); d.update(held); out.append(d)
                self.send(*json_response(out)); return
            if path.startswith('/api/held-bills/resume/'):
                hid=int(path.rsplit('/',1)[1]); r=conn.execute('SELECT * FROM held_bills WHERE id=?',(hid,)).fetchone()
                if not r: self.send(*json_response({'error':'Held bill not found.'},404)); return
                held=decode_held_bill(r['items_json']); self.send(*json_response({'ok':True,'bill':dict(r),'items':held['items'],'discount':held['discount'],'customer_phone':held['customer_phone']})); return
            if path=='/api/purchases':
                rows=conn.execute('SELECT p.*, s.name supplier FROM purchases p LEFT JOIN suppliers s ON s.id=p.supplier_id ORDER BY p.id DESC LIMIT 200').fetchall(); self.send(*json_response(serialize_rows(rows))); return
            if path=='/api/expenses':
                start=qs.get('start',[today()])[0]; end=qs.get('end',[today()])[0]; rows=conn.execute('SELECT * FROM expenses WHERE expense_date BETWEEN ? AND ? ORDER BY expense_date DESC,id DESC',(start,end)).fetchall(); self.send(*json_response(serialize_rows(rows))); return
            if path=='/api/inventory/ledger':
                pid=int(qs.get('product_id',['0'])[0]); start=qs.get('start',['1900-01-01'])[0]; end=qs.get('end',[today()])[0]
                rows=conn.execute('SELECT m.*,p.name,p.sku FROM inventory_movements m JOIN products p ON p.id=m.product_id WHERE m.product_id=? AND date(m.created_at) BETWEEN ? AND ? ORDER BY m.created_at,id',(pid,start,end)).fetchall(); self.send(*json_response(serialize_rows(rows))); return
            if path=='/api/backup':
                data=json.dumps(build_backup(conn),indent=2,default=str).encode(); fn=f'fireworks-backup-{dt.datetime.now().strftime("%Y%m%d-%H%M%S")}.json'; self.send(200,'application/json',data,{'Content-Disposition':f'attachment; filename="{fn}"'}); return
            if path=='/api/export/sales.xlsx':
                fp=make_sales_xlsx(conn); data=read_and_remove(fp); self.send(200,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',data,{'Content-Disposition':f'attachment; filename="{os.path.basename(fp)}"'}); return
            if path=='/api/export/sales-detailed.xlsx':
                fp=make_detailed_sales_xlsx(conn); data=read_and_remove(fp); self.send(200,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',data,{'Content-Disposition':f'attachment; filename="{os.path.basename(fp)}"'}); return
            if path=='/api/export/sales.pdf':
                fp=make_sales_pdf(conn); data=read_and_remove(fp); self.send(200,'application/pdf',data,{'Content-Disposition':f'attachment; filename="{os.path.basename(fp)}"'}); return
            if path.startswith('/api/invoice-pdf/'):
                iid=int(path.rsplit('/',1)[1]); fp=make_invoice_pdf(conn,iid); data=open(fp,'rb').read(); self.send(200,'application/pdf',data,{'Content-Disposition':f'inline; filename="{os.path.basename(fp)}"'}); return
            self.send(*json_response({'error':'Not found'},404))
        except Exception as e:
            traceback.print_exc(); self.send(*json_response({'error':'Something went wrong. Please try again.'},500))
        finally: conn.close()

    def do_POST(self):
        path=urllib.parse.urlparse(self.path).path
        if not self.same_origin():
            self.send(*json_response({'error':'Cross-site request rejected.'},403)); return
        if path=='/api/login':
            try: b=parse_body(self)
            except Exception: self.send(*json_response({'error':'Invalid login request.'},400)); return
            if not isinstance(b,dict): self.send(*json_response({'error':'Invalid login request.'},400)); return
            username=str(b.get('username','')).strip().lower()
            if login_blocked(username):
                self.send(*json_response({'error':'Too many failed login attempts. Try again in 15 minutes.'},429)); return
            conn=db(); u=conn.execute('SELECT * FROM users WHERE username=? AND active=1',(b.get('username',''),)).fetchone(); ok=bool(u and check_password(str(b.get('password','')),u['password_hash'])); conn.close()
            if not ok:
                record_login_failure(username)
                self.send(*json_response({'error':'Invalid username or password'},401)); return
            _login_failures.pop(username,None)
            tok=sign_token({'uid':u['id'],'sv':u['session_version'],'exp':int(dt.datetime.now().timestamp())+86400*7}); secure='; Secure' if os.getenv('FIREWORKS_SECURE_COOKIES') == '1' else ''; self.send(*json_response({'ok':True,'user':safe_user(u)}), extra_headers={'Set-Cookie':f'{COOKIE}={urllib.parse.quote(tok)}; HttpOnly; SameSite=Lax; Path=/{secure}'}); return
        if path=='/api/logout':
            secure='; Secure' if os.getenv('FIREWORKS_SECURE_COOKIES') == '1' else ''; self.send(*json_response({'ok':True}), extra_headers={'Set-Cookie':f'{COOKIE}=; Max-Age=0; Path=/; HttpOnly; SameSite=Lax{secure}'}); return
        u=self.require_auth()
        if not u:return
        if not is_admin_user(u) and not staff_api_allowed('POST', path):
            self.send(*json_response({'error':'This account does not have access to that feature.'},403)); return
        conn=db()
        try:
            if path=='/api/users':
                if not is_admin_user(u):
                    self.send(*json_response({'error':'Administrator access required.'},403)); return
                b=parse_body(self)
                username=str(b.get('username','')).strip()
                password=b.get('password','')
                if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.-]{2,31}',username):
                    raise ValueError('Username must be 3–32 characters and use only letters, numbers, dots, dashes, or underscores.')
                if not isinstance(password,str) or len(password)<8:
                    raise ValueError('Password must be at least 8 characters.')
                if conn.execute('SELECT 1 FROM users WHERE username=?',(username,)).fetchone():
                    raise ValueError('That username is already in use.')
                conn.execute('INSERT INTO users(username,password_hash,role,active,created_at) VALUES(?,?,?,1,?)',
                    (username,hash_password(password),'STAFF',now_iso()))
                user_id=conn.execute('SELECT last_insert_rowid()').fetchone()[0]
                audit(conn,u['id'],'Staff user created','USER',user_id,new={'username':username,'role':'STAFF'})
                conn.commit()
                self.send(*json_response({'ok':True,'id':user_id})); return
            if re.fullmatch(r'/api/users/\d+/status',path):
                if not is_admin_user(u):
                    self.send(*json_response({'error':'Administrator access required.'},403)); return
                user_id=int(path.split('/')[-2]); b=parse_body(self)
                target=conn.execute('SELECT id,username,role,active FROM users WHERE id=?',(user_id,)).fetchone()
                if not target:
                    self.send(*json_response({'error':'User not found.'},404)); return
                if 'active' not in b:
                    raise ValueError('Active status is required.')
                if isinstance(b['active'],int) and not isinstance(b['active'],bool) and b['active'] not in (0,1):
                    raise ValueError('Active status must be true or false.')
                if not isinstance(b['active'],(bool,int)) and str(b['active']).strip().lower() not in ('0','1','true','false','yes','no','on','off'):
                    raise ValueError('Active status must be true or false.')
                active=1 if parse_bool(b.get('active')) else 0
                if target['id']==u['id'] and not active:
                    raise ValueError('You cannot deactivate your own account.')
                if target['role']=='OWNER' and not active:
                    owners=conn.execute('SELECT COUNT(*) FROM users WHERE role="OWNER" AND active=1').fetchone()[0]
                    if owners<=1:
                        raise ValueError('At least one active administrator must remain.')
                conn.execute('UPDATE users SET active=?,session_version=session_version+? WHERE id=?',(active,0 if active else 1,user_id))
                audit(conn,u['id'],'User activated' if active else 'User deactivated','USER',user_id,
                    old={'username':target['username'],'active':target['active']},
                    new={'username':target['username'],'active':active})
                conn.commit()
                self.send(*json_response({'ok':True,'active':bool(active)})); return
            if path=='/api/import/products':
                fields,files=parse_multipart(self); fi=files.get('file')
                if not fi: raise ValueError('Attach a CSV or Excel file.')
                rawrows=csv_or_xlsx_rows(fi); headers=list(rawrows[0].keys()) if rawrows else []; mapping=json.loads(fields.get('mapping','{}')) if fields.get('mapping') else suggest_map(headers)
                required=['sku','name','unit_price']; missing=[x for x in required if not mapping.get(x)]
                if missing: raise ValueError('Mandatory columns missing: '+', '.join(missing))
                # If just preview requested
                if fields.get('mode')=='preview':
                    preview=[]; errors=[]; seen=set()
                    for ix,r in enumerate(rawrows[:100],1):
                        sku=str(r.get(mapping['sku'],'')).strip(); name=str(r.get(mapping['name'],'')).strip(); price=num(r.get(mapping['unit_price']),-1)
                        if not sku or not name or price<0: errors.append({'row':ix,'error':'Product identifier, Product Name, and a non-negative Unit Price are required'}); continue
                        if sku in seen: errors.append({'row':ix,'error':f'Duplicate Product ID: {sku}'})
                        elif conn.execute('SELECT 1 FROM products WHERE sku=?',(sku,)).fetchone(): errors.append({'row':ix,'error':f'Duplicate Product ID already exists: {sku}'})
                        seen.add(sku)
                        preview.append({k:r.get(mapping.get(k)) if mapping.get(k) else None for k in COLUMN_ALIASES})
                    self.send(*json_response({'count':len(rawrows),'headers':headers,'mapping':mapping,'preview':preview,'errors':errors})); return
                conn.execute('BEGIN'); imported=0; added=updated=0; duplicates=0; errors=[]; seen=set()
                for ix,r in enumerate(rawrows,1):
                    try:
                        sku=str(r.get(mapping['sku'],'')).strip(); name=str(r.get(mapping['name'],'')).strip(); unit=str(r.get(mapping['unit']) or 'unit').strip(); unit_price=num(r.get(mapping['unit_price']),-1); opening=num(r.get(mapping.get('opening_stock')),0)
                        if not sku or not name: raise ValueError('Product identifier and name are required')
                        if unit_price<0: raise ValueError('Unit price is required and cannot be negative')
                        if opening<0: raise ValueError('Opening stock cannot be negative')
                        if sku in seen: duplicates+=1; raise ValueError(f'Duplicate Product ID: {sku}')
                        seen.add(sku)
                        old=conn.execute('SELECT * FROM products WHERE sku=?',(sku,)).fetchone()
                        if old: duplicates+=1; raise ValueError(f'Duplicate Product ID already exists: {sku}')
                        vals=dict(category=str(r.get(mapping.get('category')) or ''),subcategory=str(r.get(mapping.get('subcategory')) or ''),brand=str(r.get(mapping.get('brand')) or ''),reorder_level=num(r.get(mapping.get('reorder_level')),0),status=str(r.get(mapping.get('status')) or 'ACTIVE').upper(),default_unit_price=unit_price)
                        if old:
                            raise ValueError(f'Duplicate Product ID already exists: {sku}')
                        else:
                            conn.execute('INSERT INTO products(sku,name,category,subcategory,unit,reorder_level,status,default_unit_price,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?)', (sku,name,vals['category'],vals['subcategory'],unit,vals['reorder_level'],vals['status'],unit_price,now_iso(),now_iso())); pid=conn.execute('SELECT last_insert_rowid()').fetchone()[0]; added+=1
                        if opening:
                            conn.execute('INSERT INTO inventory_movements(product_id,movement_type,qty_base,unit_cost,ref_type,reason,user_id,created_at) VALUES(?,?,?,?,?,?,?,?)', (pid,'OPENING',opening,'','IMPORT','Product master opening stock',u['id'],now_iso()))
                        imported+=1
                    except Exception as ex: errors.append({'row':ix,'error':str(ex)})
                if errors: conn.rollback(); self.send(*json_response({'ok':False,'error':f'Product Master validation failed for {len(errors)} row(s).','total_rows':len(rawrows),'imported':0,'failed_rows':len(errors),'duplicate_rows':duplicates,'validation_errors':errors},400)); return
                conn.commit(); audit(conn,u['id'],'Product master imported','PRODUCT',None,new={'rows':imported,'added':added,'updated':updated}); conn.commit(); self.send(*json_response({'ok':True,'total_rows':len(rawrows),'imported':imported,'added':added,'updated':updated,'failed_rows':0,'duplicate_rows':duplicates,'validation_errors':[]})); return
            if path=='/api/products':
                b=parse_body(self); sku=str(b.get('sku','')).strip(); name=str(b.get('name','')).strip(); unit_price=num(b.get('unit_price'),-1); opening=num(b.get('opening_stock'),0)
                if not sku or not name: raise ValueError('Product identifier and name are required.')
                if unit_price<0: raise ValueError('A non-negative unit price is required.')
                if opening<0: raise ValueError('Opening stock cannot be negative.')
                conn.execute('INSERT INTO products(sku,name,category,subcategory,unit,units_per_box,boxes_per_carton,units_per_carton,brand,hsn,gst_rate,cess_rate,reorder_level,status,cost_price,mrp,retail_selling_price,wholesale_selling_price,barcode,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)', (sku,name,b.get('category'),b.get('subcategory'),b.get('unit') or 'piece',num(b.get('units_per_box'),1),num(b.get('boxes_per_carton'),1),num(b.get('units_per_carton'),1),b.get('brand'),b.get('hsn'),num(b.get('gst_rate')),num(b.get('cess_rate')),num(b.get('reorder_level'),0),b.get('status','ACTIVE'),num(b.get('cost_price')),num(b.get('mrp')),num(b.get('retail_price')),num(b.get('wholesale_price')),b.get('barcode'),now_iso(),now_iso()))
                pid=conn.execute('SELECT last_insert_rowid()').fetchone()[0]; conn.execute('INSERT INTO price_versions(product_id,effective_date,mrp,retail_price,wholesale_price,cost_price,created_by,created_at) VALUES(?,?,?,?,?,?,?,?)',(pid,today(),num(b.get('mrp')),num(b.get('retail_price')),num(b.get('wholesale_price')),num(b.get('cost_price')),u['id'],now_iso())); opening=num(b.get('opening_stock')); p=conn.execute('SELECT * FROM products WHERE id=?',(pid,)).fetchone()
                conn.execute('UPDATE products SET default_unit_price=? WHERE id=?',(unit_price,pid))
                if opening:
                    qb=opening; conn.execute('INSERT INTO inventory_movements(product_id,movement_type,qty_base,unit_cost,ref_type,reason,user_id,created_at) VALUES(?,?,?,?,?,?,?,?)',(pid,'OPENING',qb,0,'MANUAL','Product master opening stock',u['id'],now_iso()))
                    if b.get('batch'): conn.execute('INSERT INTO stock_batches(product_id,batch_number,qty_base,cost,purchase_date,created_at) VALUES(?,?,?,?,?,?)',(pid,str(b.get('batch')),qb,num(b.get('cost_price')),today(),now_iso()))
                audit(conn,u['id'],'Product created','PRODUCT',pid,new=b); conn.commit(); self.send(*json_response({'ok':True,'id':pid})); return
            if path.startswith('/api/products/') and path.endswith('/remove'):
                pid=int(path.split('/')[-2]); p=conn.execute('SELECT * FROM products WHERE id=?',(pid,)).fetchone()
                if not p: raise ValueError('Product not found.')
                refs=conn.execute('SELECT (SELECT COUNT(*) FROM invoice_items WHERE product_id=?) + (SELECT COUNT(*) FROM purchase_items WHERE product_id=?) + (SELECT COUNT(*) FROM return_items WHERE product_id=?) c',(pid,pid,pid)).fetchone()['c']
                if refs:
                    conn.execute('UPDATE products SET status="INACTIVE",updated_at=? WHERE id=?',(now_iso(),pid)); action='Product deactivated'
                else:
                    conn.execute('DELETE FROM stock_batches WHERE product_id=?',(pid,)); conn.execute('DELETE FROM inventory_movements WHERE product_id=?',(pid,)); conn.execute('DELETE FROM price_versions WHERE product_id=?',(pid,)); conn.execute('DELETE FROM products WHERE id=?',(pid,)); action='Product removed'
                audit(conn,u['id'],action,'PRODUCT',pid,old=dict(p)); conn.commit(); self.send(*json_response({'ok':True,'action':action.lower()})); return
            if path=='/api/customers':
                b=parse_body(self); code=b.get('customer_code') or f'CUST-{secrets.randbelow(99999):05d}'
                conn.execute('INSERT INTO customers(customer_code,name,phone,address,gstin,customer_type,credit_enabled,credit_limit,state_code,notes,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(code,b.get('name'),b.get('phone'),b.get('address'),b.get('gstin'),b.get('customer_type','RETAIL'),1 if b.get('credit_enabled') else 0,num(b.get('credit_limit')),b.get('state_code') or settings(conn).get('state_code'),b.get('notes'),now_iso(),now_iso())); cid=conn.execute('SELECT last_insert_rowid()').fetchone()[0]; audit(conn,u['id'],'Customer created','CUSTOMER',cid,new=b); conn.commit(); self.send(*json_response({'ok':True,'id':cid})); return
            if path=='/api/sales':
                b=parse_body(self); conn.execute('BEGIN')
                iid=create_invoice(conn,b,u['id']); conn.commit(); self.send(*json_response({'ok':True,'invoice_id':iid,'invoice_number':conn.execute('SELECT invoice_number FROM invoices WHERE id=?',(iid,)).fetchone()['invoice_number']})); return
            if re.fullmatch(r'/api/invoices/\d+/edit',path):
                invoice_id=int(path.split('/')[-2]); b=parse_body(self); conn.execute('BEGIN IMMEDIATE'); result=edit_invoice(conn,invoice_id,b,u['id']); conn.commit(); self.send(*json_response({'ok':True,**result})); return
            if re.fullmatch(r'/api/invoices/\d+/void',path):
                invoice_id=int(path.split('/')[-2]); conn.execute('BEGIN'); result=void_invoice(conn,invoice_id,u['id']); conn.commit(); self.send(*json_response({'ok':True,**result})); return
            if path=='/api/held-bills':
                b=parse_body(self); ref=f'HOLD-{dt.datetime.now().strftime("%Y%m%d%H%M%S")}-{secrets.randbelow(999):03d}'
                conn.execute('INSERT INTO held_bills(ref,customer_id,customer_name,pricing_mode,items_json,notes,created_at,user_id) VALUES(?,?,?,?,?,?,?,?)',(ref,b.get('customer_id'),b.get('customer_name'),b.get('pricing_mode','STANDARD'),encode_held_bill(b),b.get('notes'),now_iso(),u['id'])); hid=conn.execute('SELECT last_insert_rowid()').fetchone()[0]; conn.commit(); self.send(*json_response({'ok':True,'id':hid,'ref':ref})); return
            if path.startswith('/api/held-bills/resume/'):
                hid=int(path.rsplit('/',1)[1]); conn.execute('BEGIN IMMEDIATE')
                try:
                    held=consume_held_bill(conn,hid); audit(conn,u['id'],'Held bill resumed','HELD_BILL',hid); conn.commit()
                except ValueError as error:
                    conn.rollback(); self.send(*json_response({'error':str(error)},404)); return
                self.send(*json_response({'ok':True,**held})); return
            if path.startswith('/api/held-bills/delete/'):
                hid=int(path.rsplit('/',1)[1]); conn.execute('DELETE FROM held_bills WHERE id=?',(hid,)); conn.commit(); self.send(*json_response({'ok':True})); return
            if path=='/api/payments':
                b=parse_body(self); iid=int(b['invoice_id']); inv=conn.execute('SELECT * FROM invoices WHERE id=?',(iid,)).fetchone(); amt=num(b.get('amount'))
                if not inv: raise ValueError('Invoice not found.')
                if amt<=0 or amt>inv['outstanding']+0.01: raise ValueError('Payment must be greater than 0 and not exceed outstanding.')
                conn.execute('INSERT INTO payments(invoice_id,customer_id,payment_date,mode,amount,reference,notes,user_id,created_at) VALUES(?,?,?,?,?,?,?,?,?)',(iid,inv['customer_id'],today(),b.get('mode','Cash'),amt,b.get('reference'),b.get('notes'),u['id'],now_iso()))
                newout=round(inv['outstanding']-amt,2); conn.execute('UPDATE invoices SET amount_received=amount_received+?,outstanding=? WHERE id=?',(amt,newout,iid)); audit(conn,u['id'],'Payment collected','PAYMENT',iid,new=b); conn.commit(); self.send(*json_response({'ok':True,'outstanding':newout})); return
            if path=='/api/purchases':
                b=parse_body(self); items=b.get('items') or []; conn.execute('BEGIN'); pnum=next_seq(conn,'purchase'); subtotal=gst=0
                conn.execute('INSERT INTO purchases(purchase_number,purchase_date,supplier_id,supplier_name,invoice_number,subtotal,gst,freight,total,user_id,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)',(pnum,b.get('purchase_date') or today(),b.get('supplier_id'),b.get('supplier_name'),b.get('invoice_number'),0,0,num(b.get('freight')),0,u['id'],now_iso())); purid=conn.execute('SELECT last_insert_rowid()').fetchone()[0]
                for it in items:
                    p=conn.execute('SELECT * FROM products WHERE id=?',(int(it['product_id']),)).fetchone(); qty=num(it.get('qty')); su=it.get('sell_unit') or p['unit']; qb=qty*stock_base_factor(p,su); cost=num(it.get('cost_price')); gr=0; ga=0; subtotal+=qty*cost
                    batchid=None
                    if b.get('batch_tracking',True):
                        conn.execute('INSERT INTO stock_batches(product_id,batch_number,supplier,manufacturer,purchase_date,qty_base,cost,location,notes,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)',(p['id'],it.get('batch_number'),b.get('supplier_name'),p['brand'],b.get('purchase_date') or today(),qb,cost,it.get('location'),it.get('notes'),now_iso())); batchid=conn.execute('SELECT last_insert_rowid()').fetchone()[0]
                    conn.execute('INSERT INTO purchase_items(purchase_id,product_id,qty_input,sell_unit,qty_base,cost_price,gst_rate,gst_amount,batch_number,batch_id) VALUES(?,?,?,?,?,?,?,?,?,?)',(purid,p['id'],qty,su,qb,cost,gr,ga,it.get('batch_number'),batchid))
                    conn.execute('INSERT INTO inventory_movements(product_id,movement_type,qty_base,unit_cost,ref_type,ref_id,reason,user_id,batch_id,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)',(p['id'],'PURCHASE',qb,cost,'PURCHASE',purid,'Purchase/GRN',u['id'],batchid,now_iso()))
                total=round(subtotal+gst+num(b.get('freight')),2); conn.execute('UPDATE purchases SET subtotal=?,gst=?,total=? WHERE id=?',(subtotal,gst,total,purid)); audit(conn,u['id'],'Stock purchase posted','PURCHASE',purid,new=b); conn.commit(); self.send(*json_response({'ok':True,'purchase_id':purid,'purchase_number':pnum,'total':total})); return
            if path=='/api/inventory/adjust':
                b=parse_body(self); pid=int(b['product_id']); qty=num(b.get('qty')); typ=b.get('type','OTHER'); delta=-abs(qty) if typ in ('DAMAGE','BREAKAGE','MOISTURE','EXPIRED','MISSING') else abs(qty)
                if delta<0 and not settings(conn).get('negative_stock_allowed',False) and current_stock(conn,pid)+delta<0: raise ValueError('Adjustment would create negative stock.')
                conn.execute('INSERT INTO inventory_movements(product_id,movement_type,qty_base,unit_cost,ref_type,reason,user_id,created_at,notes) VALUES(?,?,?,?,?,?,?,?,?)',(pid,'ADJUSTMENT',delta,num(b.get('unit_cost')), 'ADJUSTMENT',b.get('reason') or typ,u['id'],now_iso(),b.get('notes'))); audit(conn,u['id'],'Stock adjusted','INVENTORY',pid,new=b); conn.commit(); self.send(*json_response({'ok':True,'stock':current_stock(conn,pid)})); return
            if path=='/api/inventory/remove':
                b=parse_body(self); pid=int(b['product_id']); qty=num(b.get('qty')); p=conn.execute('SELECT * FROM products WHERE id=?',(pid,)).fetchone()
                if not p: raise ValueError('Product not found.')
                if qty<=0: raise ValueError('Removal quantity must be greater than 0.')
                su=b.get('unit') or p['unit']; delta=-qty*stock_base_factor(p,su)
                if not settings(conn).get('negative_stock_allowed',False) and current_stock(conn,pid)+delta<0: raise ValueError('Removal would create negative stock.')
                conn.execute('INSERT INTO inventory_movements(product_id,movement_type,qty_base,unit_cost,ref_type,reason,user_id,created_at,notes) VALUES(?,?,?,?,?,?,?,?,?)',(pid,'STOCK_REMOVAL',delta,num(b.get('unit_cost')), 'INVENTORY_REMOVAL',b.get('reason') or 'Stock removed',u['id'],now_iso(),b.get('notes')))
                audit(conn,u['id'],'Stock removed','INVENTORY',pid,new=b); conn.commit(); self.send(*json_response({'ok':True,'stock':current_stock(conn,pid)})); return
            if path=='/api/returns':
                b=parse_body(self); iid=int(b['invoice_id']); inv=conn.execute('SELECT * FROM invoices WHERE id=? AND status="FINALIZED"',(iid,)).fetchone();
                if not inv: raise ValueError('Finalized invoice not found.')
                conn.execute('BEGIN'); rnum=next_seq(conn,'return'); rows=[]; refund=0
                for it in b.get('items') or []:
                    ii=conn.execute('SELECT * FROM invoice_items WHERE id=? AND invoice_id=?',(int(it['invoice_item_id']),iid)).fetchone(); qty=num(it.get('qty'))
                    already=conn.execute('SELECT COALESCE(SUM(qty_input),0) q FROM return_items ri JOIN returns r ON r.id=ri.return_id WHERE ri.invoice_item_id=? AND r.status="POSTED"',(ii['id'],)).fetchone()['q']
                    if not ii or qty<=0 or already+qty>ii['qty_input']+0.0001: raise ValueError('Invalid return quantity for one of the items.')
                    qb=ii['qty_base']/ii['qty_input']*qty; ref=ii['unit_price']*qty-ii['discount']*(qty/ii['qty_input']); rows.append((ii,qty,qb,ref)); refund+=ref
                if not rows: raise ValueError('Select at least one item to return.')
                conn.execute('INSERT INTO returns(return_number,invoice_id,customer_id,return_date,reason,total_refund,user_id,created_at) VALUES(?,?,?,?,?,?,?,?)',(rnum,iid,inv['customer_id'],today(),b.get('reason'),refund,u['id'],now_iso())); rid=conn.execute('SELECT last_insert_rowid()').fetchone()[0]
                for ii,qty,qb,ref in rows:
                    conn.execute('INSERT INTO return_items(return_id,invoice_item_id,product_id,qty_input,qty_base,refund_amount) VALUES(?,?,?,?,?,?)',(rid,ii['id'],ii['product_id'],qty,qb,ref))
                    conn.execute('INSERT INTO inventory_movements(product_id,movement_type,qty_base,unit_cost,ref_type,ref_id,reason,user_id,created_at) VALUES(?,?,?,?,?,?,?,?,?)',(ii['product_id'],'RETURN_IN',qb,ii['cost_snapshot'],'RETURN',rid,'Sales return',u['id'],now_iso()))
                audit(conn,u['id'],'Sale returned','RETURN',rid,new={'invoice_id':iid,'refund':refund}); conn.commit(); self.send(*json_response({'ok':True,'return_number':rnum,'refund':refund})); return
            if path=='/api/expenses':
                b=parse_body(self); amt=num(b.get('amount')); 
                if amt<=0: raise ValueError('Expense amount must be greater than 0.')
                expense_date=clean_date(b.get('expense_date'),'Expense date') or today()
                conn.execute('INSERT INTO expenses(expense_date,category,amount,payment_mode,vendor,notes,attachment_name,recurring,user_id,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)',(expense_date,b.get('category'),amt,b.get('payment_mode','Cash'),b.get('vendor'),b.get('notes'),b.get('attachment_name'),1 if parse_bool(b.get('recurring')) else 0,u['id'],now_iso())); eid=conn.execute('SELECT last_insert_rowid()').fetchone()[0]; audit(conn,u['id'],'Expense added','EXPENSE',eid,new=b); conn.commit(); self.send(*json_response({'ok':True,'id':eid})); return
            if path=='/api/recurring-expenses':
                b=parse_body(self); amount=num(b.get('amount'))
                if amount<=0: raise ValueError('Recurring expense amount must be greater than 0.')
                if not b.get('category') or not b.get('frequency') or not b.get('start_date'): raise ValueError('Category, frequency, and start date are required.')
                start_date=clean_date(b.get('start_date'),'Start date',required=True); end_date=clean_date(b.get('end_date'),'End date')
                conn.execute('INSERT INTO recurring_expenses(category,amount,payment_mode,vendor,frequency,start_date,end_date,active,notes,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)',(b.get('category'),amount,b.get('payment_mode','Cash'),b.get('vendor'),b.get('frequency'),start_date,end_date,1,b.get('notes'),now_iso()))
                rid=conn.execute('SELECT last_insert_rowid()').fetchone()[0]; audit(conn,u['id'],'Recurring expense added','RECURRING_EXPENSE',rid,new=b); conn.commit(); self.send(*json_response({'ok':True,'id':rid})); return
            if path=='/api/pricing/bulk':
                b=parse_body(self); products=b.get('product_ids') or []; pct=float(b.get('pct',0)); field=b.get('field','retail_price'); eff=b.get('effective_date') or today(); conn.execute('BEGIN')
                if field not in ('retail_price','wholesale_price','mrp','cost_price'): raise ValueError('Invalid price field.')
                for pid in products:
                    p=get_price(conn,int(pid),eff)
                    if not p: continue
                    vals=dict(mrp=p['mrp'],retail_price=p['retail_price'],wholesale_price=p['wholesale_price'],cost_price=p['cost_price']); vals[field]=money(vals[field]*(1+pct/100))
                    conn.execute('INSERT INTO price_versions(product_id,effective_date,mrp,retail_price,wholesale_price,cost_price,created_by,created_at) VALUES(?,?,?,?,?,?,?,?)',(pid,eff,vals['mrp'],vals['retail_price'],vals['wholesale_price'],vals['cost_price'],u['id'],now_iso())); conn.execute('UPDATE products SET mrp=?,retail_selling_price=?,wholesale_selling_price=?,cost_price=?,updated_at=? WHERE id=?',(vals['mrp'],vals['retail_price'],vals['wholesale_price'],vals['cost_price'],now_iso(),int(pid))); audit(conn,u['id'],'Bulk price update','PRICING',int(pid),old=p,new=vals)
                conn.commit(); self.send(*json_response({'ok':True,'updated':len(products)})); return
            if path=='/api/pricing':
                b=parse_body(self); pid=int(b['product_id']); unit_price=num(b.get('unit_price'),-1)
                if unit_price<0: raise ValueError('Unit price cannot be negative.')
                p=conn.execute('SELECT * FROM products WHERE id=?',(pid,)).fetchone()
                if not p: raise ValueError('Product not found.')
                conn.execute('UPDATE products SET default_unit_price=?,updated_at=? WHERE id=?',(unit_price,now_iso(),pid)); audit(conn,u['id'],'Default unit price changed','PRODUCT',pid,old={'default_unit_price':p['default_unit_price']},new={'default_unit_price':unit_price}); conn.commit(); self.send(*json_response({'ok':True})); return
            if path=='/api/password':
                b=parse_body(self); oldp=b.get('current_password',''); newp=b.get('new_password','')
                if len(newp)<8: raise ValueError('New password must be at least 8 characters.')
                urow=conn.execute('SELECT * FROM users WHERE id=?',(u['id'],)).fetchone()
                if not check_password(oldp,urow['password_hash']): raise ValueError('Current password is incorrect.')
                sv=urow['session_version']+1
                conn.execute('UPDATE users SET password_hash=?,session_version=? WHERE id=?',(hash_password(newp),sv,u['id'])); audit(conn,u['id'],'Password changed','USER',u['id']); conn.commit()
                tok=sign_token({'uid':u['id'],'sv':sv,'exp':int(dt.datetime.now().timestamp())+86400*7}); secure='; Secure' if os.getenv('FIREWORKS_SECURE_COOKIES') == '1' else ''
                self.send(*json_response({'ok':True}), extra_headers={'Set-Cookie':f'{COOKIE}={urllib.parse.quote(tok)}; HttpOnly; SameSite=Lax; Path=/{secure}'}); return
            if path=='/api/settings':
                b=parse_body(self); st=b.get('settings') or {}; tg=b.get('targets') or {}; old=settings(conn); conn.execute('UPDATE business_settings SET data_json=? WHERE id=1',(json.dumps(st),)); conn.execute('UPDATE targets SET daily=?,weekly=?,monthly=?,seasonal=?,season_start=?,season_end=? WHERE id=1',(num(tg.get('daily')),num(tg.get('weekly')),num(tg.get('monthly')),num(tg.get('seasonal')),tg.get('season_start'),tg.get('season_end'))); audit(conn,u['id'],'Settings updated','SETTINGS',1,old={'settings':old},new=b); conn.commit(); self.send(*json_response({'ok':True})); return
            self.send(*json_response({'error':'Not found'},404))
        except Exception as e:
            conn.rollback(); traceback.print_exc()
            message=str(e) if isinstance(e,ValueError) and not isinstance(e,json.JSONDecodeError) else ('That record conflicts with an existing one (duplicate value).' if isinstance(e,sqlite3.IntegrityError) else 'The request could not be processed.')
            self.send(*json_response({'error':message},400))
        finally: conn.close()


def sales_kpis(conn,start,end):
    r=conn.execute('SELECT COALESCE(SUM(grand_total),0) sales, COUNT(*) bills, COALESCE(AVG(grand_total),0) avg_bill, COALESCE(SUM(discount),0) discounts, COALESCE(SUM(cgst+sgst+igst+cess),0) tax, COALESCE(SUM(amount_received),0) received, COALESCE(SUM(outstanding),0) credit FROM invoices WHERE invoice_date BETWEEN ? AND ? AND status="FINALIZED"',(start,end)).fetchone()
    c=conn.execute('SELECT COALESCE(SUM(ii.cost_snapshot*ii.qty_input),0) cogs FROM invoice_items ii JOIN invoices i ON i.id=ii.invoice_id WHERE i.invoice_date BETWEEN ? AND ? AND i.status="FINALIZED"',(start,end)).fetchone()['cogs']
    gp=r['sales']-c; gm=(gp/r['sales']*100) if r['sales'] else 0
    return {'sales':money(r['sales']),'bills':r['bills'],'avg_bill':money(r['avg_bill']),'discounts':money(r['discounts']),'cogs':money(c),'gross_profit':money(gp),'gross_margin':money(gm)}

def expense_total(conn,start,end): return money(conn.execute('SELECT COALESCE(SUM(amount),0) a FROM expenses WHERE expense_date BETWEEN ? AND ?',(start,end)).fetchone()['a'])

def resolve_preset(p):
    d=dt.date.today()
    if p=='today': return d.isoformat(),d.isoformat()
    if p=='yesterday': x=d-dt.timedelta(days=1); return x.isoformat(),x.isoformat()
    if p=='7d': return (d-dt.timedelta(days=6)).isoformat(),d.isoformat()
    if p=='month': return d.replace(day=1).isoformat(),d.isoformat()
    if p=='season':
        y=d.year; return f'{y}-10-01',f'{y}-10-31'
    return d.isoformat(),d.isoformat()

def read_and_remove(fp):
    data=open(fp,'rb').read()
    try: os.remove(fp)
    except OSError: pass
    return data

def make_sales_xlsx(conn):
    wb=Workbook(); ws=wb.active; ws.title='Sales'
    headers=['Invoice','Date','Customer','Customer Type','Pricing Mode','Subtotal','Discount','Total','Received','Billed By','Edited By','Edited At','Revised From']
    ws.append(headers)
    for r in conn.execute(INVOICE_SELECT+' WHERE i.status="FINALIZED" ORDER BY i.invoice_date,i.id'):
        ws.append([excel_safe(v) for v in [r['invoice_number'],r['invoice_date'],r['customer_name_snapshot'],r['customer_type_snapshot'],r['pricing_mode'],r['subtotal'],r['discount'],r['grand_total'],r['amount_received'],r['billed_by_name'],r['edited_by_name'],r['edited_at'],r['revised_from_number']]])
    fp=os.path.join(EXPORT_DIR,f'sales-{dt.datetime.now().strftime("%Y%m%d-%H%M%S")}.xlsx'); wb.save(fp); return fp

def make_detailed_sales_xlsx(conn):
    wb=Workbook(); sales=wb.active; sales.title='Sales'
    sales.append(['Invoice','Date','Customer','Customer Type','Pricing Mode','Subtotal','Discount','Total','Received','Outstanding','Status','Notes','Billed By','Edited By','Edited At','Revised From','Edit Reason'])
    sales_rows=conn.execute(INVOICE_SELECT+' WHERE i.status="FINALIZED" ORDER BY i.invoice_date,i.id').fetchall()
    for inv in sales_rows:
        sales.append([excel_safe(v) for v in [inv['invoice_number'],inv['invoice_date'],inv['customer_name_snapshot'],inv['customer_type_snapshot'],inv['pricing_mode'],inv['subtotal'],inv['discount'],inv['grand_total'],inv['amount_received'],inv['outstanding'],inv['status'],inv['notes'],inv['billed_by_name'],inv['edited_by_name'],inv['edited_at'],inv['revised_from_number'],inv['revision_reason']]])
    items=wb.create_sheet('Sale Items')
    items.append(['Invoice','Date','Customer','SKU','Product','Quantity','Unit','Base Quantity','Unit Price','Discount','Line Total','Cost Total','Gross Profit'])
    for row in conn.execute('SELECT i.invoice_number,i.invoice_date,i.customer_name_snapshot,ii.sku_snapshot,ii.product_name_snapshot,ii.qty_input,ii.sell_unit,ii.qty_base,ii.unit_price,ii.discount,ii.taxable_amount,ii.cost_snapshot FROM invoice_items ii JOIN invoices i ON i.id=ii.invoice_id WHERE i.status="FINALIZED" ORDER BY i.invoice_date,i.id,ii.id'):
        line_total=money(row['unit_price']*row['qty_input']-row['discount'])
        cost_total=money(row['cost_snapshot']*row['qty_input'])
        items.append([excel_safe(v) for v in [row['invoice_number'],row['invoice_date'],row['customer_name_snapshot'],row['sku_snapshot'],row['product_name_snapshot'],row['qty_input'],row['sell_unit'],row['qty_base'],row['unit_price'],row['discount'],line_total,cost_total,money(line_total-cost_total)]])
    payments=wb.create_sheet('Payments')
    payments.append(['Invoice','Sale Date','Customer','Payment Date','Mode','Amount','Reference','Notes'])
    for row in conn.execute('SELECT i.invoice_number,i.invoice_date,i.customer_name_snapshot,p.payment_date,p.mode,p.amount,p.reference,p.notes FROM payments p JOIN invoices i ON i.id=p.invoice_id WHERE i.status="FINALIZED" ORDER BY i.invoice_date,i.id,p.id'):
        payments.append([excel_safe(v) for v in row])
    for sheet in wb.worksheets:
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        for column in sheet.columns:
            letter=column[0].column_letter
            sheet.column_dimensions[letter].width=min(max(max(len(str(cell.value or '')) for cell in column)+2,12),36)
    fp=os.path.join(EXPORT_DIR,f'sales-detailed-{dt.datetime.now().strftime("%Y%m%d-%H%M%S")}.xlsx'); wb.save(fp); return fp

def make_sales_pdf(conn):
    fp=os.path.join(EXPORT_DIR,f'sales-{dt.datetime.now().strftime("%Y%m%d-%H%M%S")}.pdf')
    c=canvas.Canvas(fp,pagesize=A4); w,h=A4; y=h-18*mm; c.setFont('Helvetica-Bold',14); c.drawString(18*mm,y,settings(conn)['business_name']+' — Sales Report'); y-=8*mm
    c.setFont('Helvetica',8)
    rows=conn.execute('SELECT invoice_number,invoice_date,customer_name_snapshot,customer_type_snapshot,grand_total FROM invoices WHERE status="FINALIZED" ORDER BY invoice_date,id').fetchall()
    for x,label in [(18,'Invoice'),(55,'Date'),(82,'Customer'),(145,'Type'),(170,'Total')]: c.drawString(x*mm,y,label)
    y-=4*mm; c.line(18*mm,y,w-18*mm,y); y-=5*mm
    for r in rows:
        c.drawString(18*mm,y,str(r['invoice_number'])[-18:]); c.drawString(55*mm,y,str(r['invoice_date'])); c.drawString(82*mm,y,str(r['customer_name_snapshot'])[:25]); c.drawString(145*mm,y,str(r['customer_type_snapshot'])[:10]); c.drawRightString(188*mm,y,f'{r["grand_total"]:,.2f}'); y-=4.5*mm
        if y<18*mm: c.showPage(); y=h-18*mm
    c.save(); return fp

def short_time(value):
    try: return dt.datetime.fromisoformat(str(value)).strftime('%d-%m-%Y %I:%M %p')
    except (TypeError, ValueError): return str(value or '')

def make_invoice_pdf(conn,iid):
    inv=conn.execute(INVOICE_SELECT+' WHERE i.id=?',(iid,)).fetchone()
    items=conn.execute('SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY id',(iid,)).fetchall()
    pays=conn.execute('SELECT * FROM payments WHERE invoice_id=? ORDER BY id',(iid,)).fetchall()
    st=settings(conn)
    fp=os.path.join(BILLS_DIR,f'{inv["invoice_number"]}.pdf')
    c=canvas.Canvas(fp,pagesize=A4); c.setPageCompression(0)
    w,h=A4

    left_x=18*mm
    right_x=(w/2)+6*mm
    col_width=(w/2)-24*mm
    top_margin=18*mm
    bottom_margin=18*mm
    line_h=5*mm
    footer_space=28*mm

    def fit_text(text,font,size,max_width):
        text=str(text or '')
        if stringWidth(text,font,size)<=max_width:
            return text
        while text and stringWidth(text+'...',font,size)>max_width:
            text=text[:-1]
        return text+'...' if text else ''

    def draw_header_at(x_offset,copy_label,continued=False):
        y=h-top_margin
        c.setFont('Helvetica-Bold',11)
        c.drawString(x_offset,y,fit_text(st['business_name'],'Helvetica-Bold',11,col_width-3*mm))
        y-=6*mm
        c.setFont('Helvetica',7)
        c.drawString(x_offset,y,fit_text(st.get('address',''),'Helvetica',7,col_width))
        y-=4*mm
        c.drawString(x_offset,y,fit_text(f'License: {st.get("license_number") or "N/A"}','Helvetica',7,col_width))
        c.setFont('Helvetica-Bold',10); c.drawRightString(x_offset+col_width, h-top_margin+2*mm, 'INVOICE')
        c.setFont('Helvetica-Bold',8); c.drawRightString(x_offset+col_width, h-top_margin-5*mm, copy_label)
        y-=8*mm
        c.setFont('Helvetica',7)
        c.drawString(x_offset,y,fit_text(f'Invoice: {inv["invoice_number"]}','Helvetica',7,col_width))
        y-=4*mm
        c.drawString(x_offset,y,fit_text(f'Date: {inv["invoice_date"]}','Helvetica',7,col_width)); y-=4*mm
        c.drawString(x_offset,y,fit_text(f'Customer: {inv["customer_name_snapshot"]} ({inv["customer_type_snapshot"]})','Helvetica',7,col_width)); y-=4*mm
        c.drawString(x_offset,y,fit_text(f'Billed by: {inv["billed_by_name"] or "-"}','Helvetica',7,col_width)); y-=4*mm
        if inv['revised_from_number']:
            c.setFont('Helvetica-Oblique',7)
            c.drawString(x_offset,y,fit_text(f'Revised from {inv["revised_from_number"]} - edited by {inv["edited_by_name"] or "-"} on {short_time(inv["edited_at"])}','Helvetica-Oblique',7,col_width)); y-=4*mm
            c.setFont('Helvetica',7)
        if inv['status']=='VOID':
            c.setFont('Helvetica-Bold',7)
            replaced=f' - replaced by {inv["revised_to_number"]}' if inv['revised_to_number'] else ''
            c.drawString(x_offset,y,fit_text(f'CANCELLED{replaced} ({inv["voided_by_name"] or "-"}, {short_time(inv["voided_at"])})','Helvetica-Bold',7,col_width)); y-=4*mm
            c.setFont('Helvetica',7)
        y-=1*mm
        if continued:
            c.setFont('Helvetica-Oblique',7); c.drawString(x_offset,y,'Continued'); y-=4*mm
        c.setFont('Helvetica-Bold',7)
        c.drawString(x_offset,y,'Product / Item')
        c.drawRightString(x_offset+38*mm,y,'Qty')
        c.drawRightString(x_offset+59*mm,y,'Unit Price')
        c.drawRightString(x_offset+col_width,y,'Line Total')
        y-=4*mm; c.line(x_offset,y,x_offset+col_width,y); y-=4*mm
        c.setFont('Helvetica',7)
        return y

    # paginate items in shared slices so both columns show same content on each page
    header_space = 71*mm
    usable_height = h - top_margin - header_space - bottom_margin - footer_space
    items_per_page = max(1, int(usable_height // line_h))

    total_items = len(items)
    index = 0
    page_num = 0
    while index < total_items or page_num==0:
        # for empty items still show a page with header and totals
        c.setFont('Helvetica',7)
        left_y = draw_header_at(left_x,'CUSTOMER COPY',continued=(page_num>0))
        right_y = draw_header_at(right_x,'SELLER COPY',continued=(page_num>0))

        slice_items = items[index:index+items_per_page]
        y_left = left_y
        y_right = right_y
        for it in slice_items:
            product_name=fit_text(f"{it['sku_snapshot']} {it['product_name_snapshot']}",'Helvetica',7,35*mm)
            c.drawString(left_x, y_left, product_name)
            c.drawRightString(left_x+38*mm, y_left, f"{it['qty_input']:g}")
            c.drawRightString(left_x+59*mm, y_left, f"{it['unit_price']:.2f}")
            c.drawRightString(left_x+col_width, y_left, f"{it['unit_price']*it['qty_input']:,.2f}")
            c.drawString(right_x, y_right, product_name)
            c.drawRightString(right_x+38*mm, y_right, f"{it['qty_input']:g}")
            c.drawRightString(right_x+59*mm, y_right, f"{it['unit_price']:.2f}")
            c.drawRightString(right_x+col_width, y_right, f"{it['unit_price']*it['qty_input']:,.2f}")
            y_left -= line_h; y_right -= line_h
        index += items_per_page

        # if this is the last page, draw totals and payments within each column
        if index >= total_items:
            # totals block for left
            tx = left_x
            ty = bottom_margin + footer_space - 8*mm
            c.setFont('Helvetica-Bold',7)
            c.line(tx, ty+10*mm, tx+col_width, ty+10*mm)
            ty -= 2*mm
            totals = [('Subtotal',inv['subtotal']),('Discount',inv['discount']),('Total',inv['grand_total']),('Received',inv['amount_received'])]
            if inv['outstanding'] > 0:
                totals.append(('Credit',inv['outstanding']))
            for label,val in totals:
                c.drawString(tx+34*mm, ty, label); c.drawRightString(tx+col_width, ty, f'Rs. {val:,.2f}'); ty -= 5*mm
            c.setFont('Helvetica',7)
            payments_str = ', '.join([f"{p['mode']} Rs. {p['amount']:,.2f}" for p in pays]) if pays else 'No payment recorded'
            c.drawString(tx, ty-2*mm, fit_text('Payment(s): '+payments_str,'Helvetica',7,col_width))
            c.drawString(tx, ty-8*mm, fit_text(st.get('invoice_footer','Thank you.'),'Helvetica',7,col_width))
            tx = right_x
            ty = bottom_margin + footer_space - 8*mm
            c.setFont('Helvetica-Bold',7)
            c.line(tx, ty+10*mm, tx+col_width, ty+10*mm)
            ty -= 2*mm
            for label,val in totals:
                c.drawString(tx+34*mm, ty, label); c.drawRightString(tx+col_width, ty, f'Rs. {val:,.2f}'); ty -= 5*mm
            c.setFont('Helvetica',7)
            c.drawString(tx, ty-2*mm, fit_text('Payment(s): '+payments_str,'Helvetica',7,col_width))
            c.drawString(tx, ty-8*mm, fit_text(st.get('invoice_footer','Thank you.'),'Helvetica',7,col_width))
        # if more items remain, advance page
        if index < total_items:
            c.showPage()
        page_num += 1
        if total_items==0:
            break
    c.save()
    return fp

def default_password_in_use():
    conn=db(); rows=conn.execute('SELECT password_hash FROM users WHERE active=1').fetchall(); conn.close()
    return any(check_password('admin123',r['password_hash']) for r in rows)


if __name__=='__main__':
    init_db()
    if not is_loopback_host(HOST) and default_password_in_use():
        raise SystemExit('An account still uses the demo password "admin123". Start the app normally (localhost), change the password in Settings, then start it on the network again.')
    print(f'Fireworks CRM running at http://{"localhost" if is_loopback_host(HOST) else HOST}:{PORT}')
    print('Login username: admin; password is configured through FIREWORKS_ADMIN_PASSWORD.' if os.getenv('FIREWORKS_ADMIN_PASSWORD') else 'Local demo login: admin / admin123 (change it in Settings)')
    ThreadingHTTPServer((HOST,PORT),Handler).serve_forever()

"""Excel sales exports, generated in memory."""
import io

from openpyxl import Workbook

from crm.domain.money import money

FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')


def excel_safe(value):
    """Stop text such as '=HYPERLINK(...)' being written as a live spreadsheet formula."""
    if isinstance(value, str) and value[:1] in FORMULA_PREFIXES:
        return "'" + value
    return value


def _append(sheet, values):
    sheet.append([excel_safe(value) for value in values])


def _to_bytes(workbook):
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _tidy(workbook):
    for sheet in workbook.worksheets:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = max(len(str(cell.value or '')) for cell in column) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(max(width, 12), 36)


def sales_summary(invoices):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Sales'
    sheet.append(['Invoice', 'Date', 'Customer', 'Customer Type', 'Pricing Mode', 'Subtotal', 'Discount', 'Total', 'Received',
                  'Billed By', 'Edited By', 'Edited At', 'Revised From'])
    for inv in invoices:
        _append(sheet, [inv['invoice_number'], inv['invoice_date'], inv['customer_name_snapshot'], inv['customer_type_snapshot'],
                        inv['pricing_mode'], inv['subtotal'], inv['discount'], inv['grand_total'], inv['amount_received'],
                        inv['billed_by_name'], inv['edited_by_name'], inv['edited_at'], inv['revised_from_number']])
    return _to_bytes(workbook)


def sales_detailed(invoices, items, payments, include_profit=True):
    workbook = Workbook()
    sales = workbook.active
    sales.title = 'Sales'
    sales.append(['Invoice', 'Date', 'Customer', 'Customer Type', 'Pricing Mode', 'Subtotal', 'Discount', 'Total', 'Received',
                  'Outstanding', 'Status', 'Notes', 'Billed By', 'Edited By', 'Edited At', 'Revised From', 'Edit Reason'])
    for inv in invoices:
        _append(sales, [inv['invoice_number'], inv['invoice_date'], inv['customer_name_snapshot'], inv['customer_type_snapshot'],
                        inv['pricing_mode'], inv['subtotal'], inv['discount'], inv['grand_total'], inv['amount_received'], inv['outstanding'],
                        inv['status'], inv['notes'], inv['billed_by_name'], inv['edited_by_name'], inv['edited_at'],
                        inv['revised_from_number'], inv['revision_reason']])
    item_sheet = workbook.create_sheet('Sale Items')
    item_sheet.append(['Invoice', 'Date', 'Customer', 'SKU', 'Product', 'Brand / Company', 'Quantity', 'Unit', 'Base Quantity', 'Unit Price', 'Discount',
                       'Line Total'] + (['Cost Total', 'Gross Profit'] if include_profit else []))
    for row in items:
        line_total = money(row['unit_price'] * row['qty_input'] - row['discount'])
        cost_total = money(row['cost_snapshot'] * row['qty_input'])
        profit_columns = [cost_total, money(line_total - cost_total)] if include_profit else []
        _append(item_sheet, [row['invoice_number'], row['invoice_date'], row['customer_name_snapshot'], row['sku_snapshot'],
                             row['product_name_snapshot'], row['brand_snapshot'], row['qty_input'], row['sell_unit'], row['qty_base'], row['unit_price'],
                             row['discount'], line_total] + profit_columns)
    payment_sheet = workbook.create_sheet('Payments')
    payment_sheet.append(['Invoice', 'Sale Date', 'Customer', 'Payment Date', 'Mode', 'Amount', 'Reference', 'Notes'])
    for row in payments:
        _append(payment_sheet, [row['invoice_number'], row['invoice_date'], row['customer_name_snapshot'], row['payment_date'],
                                row['mode'], row['amount'], row['reference'], row['notes']])
    _tidy(workbook)
    return _to_bytes(workbook)

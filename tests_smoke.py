import os, re, sqlite3, tempfile
import server

# Use a fresh temp DB for logic smoke tests.
fd,path=tempfile.mkstemp(suffix='.db'); os.close(fd)
old=server.DB_PATH; old_secret=server.SECRET_PATH; old_exports=server.EXPORT_DIR; old_bills=server.BILLS_DIR; server.DB_PATH=path; server.SECRET_PATH=path+'.secret'; server.EXPORT_DIR=path+'.exports'; server.BILLS_DIR=path+'.bills'; os.makedirs(server.EXPORT_DIR); os.makedirs(server.BILLS_DIR)
try:
    server.init_db()
    assert server.is_admin_user({'role':'OWNER'})
    assert not server.is_admin_user({'role':'STAFF'})
    for method, path in (
        ('GET','/api/dashboard'), ('GET','/api/analytics'), ('GET','/api/products'),
        ('GET','/api/expenses'), ('GET','/api/backup'), ('POST','/api/sales'),
        ('POST','/api/expenses'), ('POST','/api/held-bills/resume/4')
    ):
        assert server.staff_api_allowed(method,path), (method,path)
    for method, path in (
        ('GET','/api/settings'), ('GET','/api/users'), ('GET','/api/customers'),
        ('GET','/api/invoices'), ('POST','/api/users'), ('POST','/api/products'),
        ('POST','/api/held-bills/delete/4')
    ):
        assert not server.staff_api_allowed(method,path), (method,path)
    conn=server.db()
    assert conn.execute('select count(*) from products').fetchone()[0] == 0
    assert conn.execute('select count(*) from invoices where status="FINALIZED"').fetchone()[0] == 0
    conn.execute('INSERT INTO held_bills(ref,customer_name,pricing_mode,items_json,created_at) VALUES(?,?,?,?,?)',('HOLD-SMOKE','Held Customer','STANDARD',server.encode_held_bill({'items':[{'product_id':1,'qty':2}],'discount':5.25,'customer_phone':'9876543210'}),server.now_iso()))
    held_id=conn.execute('SELECT last_insert_rowid()').fetchone()[0]
    consumed=server.consume_held_bill(conn,held_id)
    assert consumed['items']==[{'product_id':1,'qty':2}] and consumed['discount']==5.25 and consumed['customer_phone']=='9876543210'
    assert conn.execute('SELECT COUNT(*) FROM held_bills WHERE id=?',(held_id,)).fetchone()[0] == 0
    conn.commit()
    assert server.stock_base_factor({'unit':'Box','units_per_box':10,'boxes_per_carton':12,'units_per_carton':120}, 'Box') == 1
    assert server.parse_bool('True') is True
    assert server.parse_bool('false') is False
    assert server.parse_bool(False) is False
    held=server.decode_held_bill(server.encode_held_bill({'items':[{'product_id':1,'qty':2}],'discount':5.25,'customer_phone':'9876543210'}))
    assert held == {'items':[{'product_id':1,'qty':2}],'discount':5.25,'customer_phone':'9876543210'}
    assert server.decode_held_bill('[{"product_id":1}]') == {'items':[{'product_id':1}],'discount':0.0,'customer_phone':''}
    assert server.expense_total(conn, server.today(), server.today()) >= 0
    user_id=conn.execute('SELECT id FROM users WHERE username="admin"').fetchone()['id']
    conn.execute('INSERT INTO products(sku,name,unit,default_unit_price,cost_price,created_at,updated_at) VALUES(?,?,?,?,?,?,?)',('TEST-1','Smoke Product','Box',20,8,server.now_iso(),server.now_iso()))
    product_id=conn.execute('SELECT last_insert_rowid()').fetchone()[0]
    conn.execute('INSERT INTO price_versions(product_id,effective_date,mrp,retail_price,wholesale_price,cost_price,created_at) VALUES(?,?,?,?,?,?,?)',(product_id,server.today(),25,20,18,8,server.now_iso()))
    assert conn.execute('SELECT COUNT(*) FROM price_versions WHERE product_id=?',(product_id,)).fetchone()[0] == 1
    conn.execute('INSERT INTO inventory_movements(product_id,movement_type,qty_base,unit_cost,reason,user_id,created_at) VALUES(?,?,?,?,?,?,?)',(product_id,'OPENING',10,8,'Smoke opening',user_id,server.now_iso()))
    conn.commit()
    invoice_id=server.create_invoice(conn,{'items':[{'product_id':product_id,'qty':2,'unit_price':20}],'payments':[{'mode':'Cash','amount':40}]},user_id)
    conn.commit()
    assert server.current_stock(conn,product_id) == 8
    summary_report=server.make_sales_xlsx(conn)
    report=server.make_detailed_sales_xlsx(conn)
    from openpyxl import load_workbook
    summary_workbook=load_workbook(summary_report,read_only=True,data_only=True)
    assert summary_workbook.sheetnames == ['Sales']
    summary_workbook.close()
    workbook=load_workbook(report,read_only=True,data_only=True)
    assert workbook.sheetnames == ['Sales','Sale Items','Payments']
    assert all(workbook[name].max_row == 2 for name in workbook.sheetnames)
    workbook.close()
    server.void_invoice(conn,invoice_id,user_id)
    conn.commit()
    assert server.current_stock(conn,product_id) == 10
    assert conn.execute('SELECT status FROM invoices WHERE id=?',(invoice_id,)).fetchone()['status'] == 'VOID'
    assert conn.execute('SELECT COUNT(*) FROM payments WHERE invoice_id=?',(invoice_id,)).fetchone()[0] == 1
    discounted_id=server.create_invoice(conn,{'customer_phone':'9876543210','items':[{'product_id':product_id,'qty':2,'unit_price':20}],'discount':5,'payments':[{'mode':'Cash','amount':35}]},user_id)
    conn.commit()
    discounted=conn.execute('SELECT subtotal,discount,taxable_amount,grand_total,amount_received FROM invoices WHERE id=?',(discounted_id,)).fetchone()
    assert tuple(discounted) == (40,5,35,35,35)
    discounted_item=conn.execute('SELECT discount,taxable_amount FROM invoice_items WHERE invoice_id=?',(discounted_id,)).fetchone()
    assert tuple(discounted_item) == (5,35)
    assert server.current_stock(conn,product_id) == 8
    invoice_pdf=server.make_invoice_pdf(conn,discounted_id)
    assert os.path.getsize(invoice_pdf) > 0
    pdf_data=open(invoice_pdf,'rb').read()
    assert len(re.findall(rb'/Type\s*/Page\b',pdf_data)) == 1
    assert b'CUSTOMER COPY' in pdf_data and b'SELLER COPY' in pdf_data
    assert b'Unit Price' in pdf_data and b'Line Total' in pdf_data
    assert b'Rs.' in pdf_data and b'\xe2\x82\xb9' not in pdf_data
    mixed_credit_id=server.create_invoice(conn,{'items':[{'product_id':product_id,'qty':1}],'payments':[{'mode':'Cash','amount':10},{'mode':'Credit','amount':10}]},user_id)
    conn.commit()
    mixed_credit=conn.execute('SELECT grand_total,amount_received,outstanding FROM invoices WHERE id=?',(mixed_credit_id,)).fetchone()
    assert tuple(mixed_credit) == (20,10,10)
    assert conn.execute('SELECT SUM(amount) FROM payments WHERE invoice_id=? AND mode="Credit"',(mixed_credit_id,)).fetchone()[0] == 10
    mixed_pdf=server.make_invoice_pdf(conn,mixed_credit_id)
    mixed_pdf_data=open(mixed_pdf,'rb').read()
    assert len(re.findall(rb'/Type\s*/Page\b',mixed_pdf_data)) == 1
    assert b'Credit' in mixed_pdf_data and b'10.00' in mixed_pdf_data
    assert b'Rs.' in mixed_pdf_data and b'\xe2\x82\xb9' not in mixed_pdf_data
    credit_only_id=server.create_invoice(conn,{'items':[{'product_id':product_id,'qty':1}],'payments':[{'mode':'Credit','amount':20}]},user_id)
    conn.commit()
    credit_only=conn.execute('SELECT grand_total,amount_received,outstanding FROM invoices WHERE id=?',(credit_only_id,)).fetchone()
    assert tuple(credit_only) == (20,0,20)
    assert conn.execute('SELECT SUM(grand_total) FROM invoices WHERE status="FINALIZED"').fetchone()[0] == 75
    try:
        server.create_invoice(conn,{'items':[{'product_id':product_id,'qty':1}],'payments':[{'mode':'Cash','amount':19}]},user_id)
        raise AssertionError('Under-allocated invoice payment was accepted.')
    except ValueError as exc:
        assert 'equal the invoice total' in str(exc)
    server.void_invoice(conn,mixed_credit_id,user_id)
    server.void_invoice(conn,credit_only_id,user_id)
    conn.commit()
    assert server.current_stock(conn,product_id) == 8
    try:
        server.create_invoice(conn,{'customer_phone':'123','items':[{'product_id':product_id,'qty':1,'unit_price':20}]},user_id)
        raise AssertionError('Invalid mobile number was accepted.')
    except ValueError as exc:
        assert '10 digits' in str(exc)
    try:
        server.create_invoice(conn,{'items':[{'product_id':product_id,'qty':1,'unit_price':20}],'discount':21},user_id)
        raise AssertionError('Discount above subtotal was accepted.')
    except ValueError as exc:
        assert 'subtotal' in str(exc)
    server.void_invoice(conn,discounted_id,user_id)
    conn.commit()
    assert server.current_stock(conn,product_id) == 10
    fixed_price_id=server.create_invoice(conn,{'items':[{'product_id':product_id,'qty':1,'unit_price':999}],'payments':[{'mode':'Cash','amount':20}]},user_id)
    conn.commit()
    fixed_price=conn.execute('SELECT unit_price FROM invoice_items WHERE invoice_id=?',(fixed_price_id,)).fetchone()['unit_price']
    assert fixed_price == 20
    server.void_invoice(conn,fixed_price_id,user_id)
    conn.commit()
    assert server.current_stock(conn,product_id) == 10
    conn.close()
    print('SMOKE TEST PASSED')
finally:
    try: os.remove(path)
    except: pass
    try: os.remove(path+'.secret')
    except: pass
    for filename in os.listdir(server.EXPORT_DIR):
        try: os.remove(os.path.join(server.EXPORT_DIR,filename))
        except: pass
    try: os.rmdir(server.EXPORT_DIR)
    except: pass
    for filename in os.listdir(server.BILLS_DIR):
        try: os.remove(os.path.join(server.BILLS_DIR,filename))
        except: pass
    try: os.rmdir(server.BILLS_DIR)
    except: pass
    server.DB_PATH=old; server.SECRET_PATH=old_secret; server.EXPORT_DIR=old_exports; server.BILLS_DIR=old_bills
